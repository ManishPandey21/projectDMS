"""Frozen baseline, EOT submission/determination, and complete-history exports."""

from __future__ import annotations

import csv
import io
from typing import Any, Dict, Iterable, List, Sequence, Tuple


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "date"):
        try:
            return value.date().isoformat()
        except Exception:
            pass
    return str(value)


def _metadata_rows(metadata: Dict[str, Any]) -> List[List[str]]:
    return [[str(key), _cell(value)] for key, value in metadata.items()]


def _csv(metadata: Dict[str, Any], headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(_metadata_rows(metadata))
    writer.writerow([])
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_cell(value) for value in row])
    return buffer.getvalue()


def _xlsx(metadata: Dict[str, Any], headers: Sequence[str], rows: Iterable[Sequence[Any]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for row in _metadata_rows(metadata):
        ws.append(row)
    ws.append([])
    ws.append(list(headers))
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(value) for value in row])
    for column in ws.columns:
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf(metadata: Dict[str, Any], headers: Sequence[str], rows: Iterable[Sequence[Any]], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    for key, value in metadata.items():
        story.append(Paragraph(f"<b>{key}:</b> {_cell(value)}", styles["BodyText"]))
    story.append(Spacer(1, 10))
    data = [list(headers)] + [[_cell(value) for value in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
    ]))
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def render(
    format: str,
    metadata: Dict[str, Any],
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    title: str,
) -> str | bytes:
    if format == "csv":
        return _csv(metadata, headers, rows)
    if format == "xlsx":
        return _xlsx(metadata, headers, rows, title)
    if format == "pdf":
        return _pdf(metadata, headers, rows, title)
    raise ValueError(f"Unsupported export format: {format}")


def baseline_table(baseline: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "Milestone Ref", "Description", "Contractual Week",
        "Original Contractual Key Date", "Responsible Party", "Remarks",
    ]
    rows = [[
        item.get("milestone_ref"),
        item.get("description") or item.get("title"),
        item.get("contractual_week_number"),
        item.get("original_contractual_date"),
        item.get("responsible_party_id"),
        item.get("remarks"),
    ] for item in baseline.get("items") or []]
    return headers, rows


def submission_table(submission: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "Milestone Ref", "Description", "Contractual Date At Submission",
        f"{submission.get('revision_label')} Submitted Date", "Claimed Extension Days", "Remarks",
    ]
    rows = [[
        item.get("milestone_ref"), item.get("description"),
        item.get("contractual_date_at_submission"), item.get("eot_submitted_date"),
        item.get("claimed_extension_days"), item.get("remarks"),
    ] for item in submission.get("items") or []]
    return headers, rows


def determination_table(determination: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "Milestone Ref", "Description", "Submitted Date", "Determined / Granted Date",
        "Claimed Extension Days", "Granted Extension Days", "Difference", "Result", "Remarks",
    ]
    rows: List[List[Any]] = []
    for item in determination.get("items") or []:
        claimed = item.get("claimed_extension_days")
        granted = item.get("granted_extension_days")
        difference = None if claimed is None or granted is None else int(claimed) - int(granted)
        rows.append([
            item.get("milestone_ref"), item.get("description"), item.get("submitted_date"),
            item.get("eot_granted_date"), claimed, granted, difference,
            item.get("determination_result"), item.get("remarks"),
        ])
    return headers, rows


def history_table(
    milestones: Sequence[Dict[str, Any]],
    submissions: Sequence[Dict[str, Any]],
    determinations: Sequence[Dict[str, Any]],
) -> Tuple[List[str], List[List[Any]]]:
    ordered = sorted(submissions, key=lambda row: int(row.get("revision_number") or 0))
    headers = ["Ref", "Description", "Original Date"]
    for submission in ordered:
        label = submission.get("revision_label") or f"EOT-{submission.get('revision_number')}"
        headers += [f"{label} Submitted", f"{label} Granted", f"{label} Status"]
    headers += ["Current Contractual Date", "Actual Achievement Date"]

    submission_items = {
        (str(submission.get("_id")), str(item.get("milestone_ref")).casefold()): item
        for submission in ordered
        for item in submission.get("items") or []
    }
    determination_by_submission_ref: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for determination in sorted(
        determinations,
        key=lambda row: row.get("frozen_at") or row.get("created_at"),
    ):
        for submission_id in determination.get("eot_submission_ids") or []:
            for item in determination.get("items") or []:
                determination_by_submission_ref[
                    (str(submission_id), str(item.get("milestone_ref")).casefold())
                ] = {**item, "determination_status": determination.get("status"), "frozen_at": determination.get("frozen_at")}

    rows: List[List[Any]] = []
    for milestone in milestones:
        ref = str(milestone.get("milestone_ref") or "")
        row: List[Any] = [
            ref,
            milestone.get("description") or milestone.get("title"),
            milestone.get("original_planned_key_date"),
        ]
        for submission in ordered:
            key = (str(submission.get("_id")), ref.casefold())
            submitted = submission_items.get(key) or {}
            determined = determination_by_submission_ref.get(key) or {}
            status = determined.get("determination_result") or (
                "pending" if submitted else "not_affected"
            )
            row += [submitted.get("eot_submitted_date"), determined.get("eot_granted_date"), status]
        row += [milestone.get("current_approved_key_date"), milestone.get("actual_achievement_date")]
        rows.append(row)
    return headers, rows
