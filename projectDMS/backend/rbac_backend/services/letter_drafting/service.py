from __future__ import annotations

import hashlib
import io
import re
import zipfile
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional

from fastapi import HTTPException

from ...models.letter import Letter
from ...models.letter_drafting import (
    ApprovalStage,
    ApprovalStep,
    ApproveStageRequest,
    AssignReviewerRequest,
    ConfirmAnalysisRequest,
    ConfirmPlanRequest,
    CyclicIterationTrace,
    DraftAssertionSupport,
    DraftAuditResponse,
    DraftArtifact,
    DraftConfidenceScores,
    DraftContextPack,
    DraftExecutionEffect,
    DraftOutboxEvent,
    DraftCommentRequest,
    DraftGovernanceResponse,
    DraftMetricBottleneck,
    DraftMetricBreakdownItem,
    DraftMetricKpi,
    DraftMetricTrendPoint,
    DraftQualityDashboardResponse,
    DraftQualityRiskItem,
    DraftingStartRequest,
    DraftingStartResponse,
    DraftMode,
    DraftRole,
    DraftRun,
    DraftRunCreateRequest,
    DraftReviewAssignment,
    DraftReviewComment,
    ExactClauseSearchRequest,
    ExactReferenceSearchRequest,
    FreezeSectionsRequest,
    FrozenDraftSection,
    IncomingLetterAnalysis,
    LegalRiskReport,
    LockParagraphsRequest,
    ProbingQuestion,
    ReviseDraftRequest,
    ReviseSectionsRequest,
    ReturnForCorrectionRequest,
    SectionRevisionRecord,
    SourceEvidence,
    SourceLedgerResponse,
    UserDirectionRequest,
    UserDirectionAnswer,
    ValidationFinding,
    ValidationReport,
)
from ...models.ai_guardrails import GuardrailReport
from ...models.evidence_ledger import EvidenceLedgerEntry
from ...models.notification import NotificationPriority, NotificationSeverity, NotificationType
from ...core.config import settings
from ...services.ai_guardrails import AIOutputGuardrailService
from ...services.authorization_service import AuthorizationService
from ...services.policy_service import PolicyService
from ...services.contract_service import ContractService
from ...services.conversation_service import ConversationService
from ...services.document_service import DocumentService
from ...models.contract_models import ContractSearchRequest
from ...core.security import build_scope_query
from ...services.letter_service import LetterService
from ...services.falkor_graph_service import normalize_letter_code
from ...services.file_object_service import FileObjectService
from ...utils.notification_service import NotificationService
from .clause_checker import ClauseCheckingAgent
from .context import DraftContextBuilder
from .generator import DraftGenerator, StrategyPlanner
from .incoming_analyzer import IncomingLetterAnalyzer
from .input_validator import DraftInputValidator
from .legal_risk_reviewer import LegalRiskReviewer
from .locked_text import locked_instruction, verify_locked_paragraphs
from .frozen_sections import (
    apply_section_edits,
    content_hash,
    freeze_sections as capture_frozen_sections,
    preserved_sections,
    protected_anchor_changes,
    section_map,
    verify_frozen_sections,
)
from .planning import PlanningSheetBuilder
from .prompts import PromptRegistry
from .repository import DraftRunRepository
from .section_editor import SECTION_EDIT_PROMPT_VERSION, ScopedSectionEditor
from .user_direction import UserDirectionAgent
from .validator import DraftValidator


# Draft run status / mode groupings referenced across the workflow. Named once
# here so a new status or mode is classified in a single place instead of by
# hand-copied set literals scattered through the service.
_AI_DRAFT_MODES = frozenset(["draft", "review"])
_FINALIZED_STATUSES = frozenset(["approved", "exported", "issued"])
_ACTIVE_STATUSES = frozenset(["completed", "needs_attention", "blocked"])
_EXPORTED_ISSUED_STATUSES = frozenset(["exported", "issued"])


class DraftRunService:
    """Coordinates v2 letter drafting runs."""

    def __init__(self, db: Any):
        self.db = db
        self.letter_service = LetterService(db)
        self.auth_service = AuthorizationService()
        self.policy_service = PolicyService(db)
        self.repository = DraftRunRepository(db)
        self.prompt_registry = PromptRegistry(db)
        self.validator = DraftValidator()
        self.input_validator = DraftInputValidator()
        self.incoming_analyzer = IncomingLetterAnalyzer()
        self.clause_checker = ClauseCheckingAgent(db)
        self.legal_risk_reviewer = LegalRiskReviewer()
        self.planning_builder = PlanningSheetBuilder()
        self.guardrails = AIOutputGuardrailService.from_settings(settings)

    async def create_run(
        self,
        letter_id: str,
        request: DraftRunCreateRequest,
        current_user: Any,
        *,
        idempotency_key: Optional[str] = None,
        request_hash: Optional[str] = None,
        engine_metadata: Optional[Dict[str, Any]] = None,
        stop_after_generation: bool = False,
    ) -> DraftRun:
        started = datetime.now(timezone.utc)
        run_id = str(uuid.uuid4())
        warnings: list[str] = []
        trace: list[dict[str, Any]] = []
        letter = await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.create",
        )
        if idempotency_key:
            existing = await self.repository.get_by_idempotency_key(letter_id, idempotency_key)
            if existing:
                if request_hash and existing.request_hash and existing.request_hash != request_hash:
                    raise HTTPException(
                        status_code=409,
                        detail="Idempotency-Key was already used with a different drafting request",
                    )
                return existing
        role = self._resolve_role(letter, request)
        recipient_focus = request.recipient_focus or getattr(letter, "strategy_recipient", None)

        # Carry the user's line of action (answers from a previous analysis run)
        # into this run's inputs so strategy/draft stages honour it.
        supplied_directions = list(request.user_direction_answers or [])
        if request.user_direction and request.user_direction.strip():
            supplied_directions.append(
                UserDirectionAnswer(
                    question_id="free_text",
                    answer=request.user_direction.strip(),
                )
            )
        directions_block = (
            UserDirectionAgent.format_directions(supplied_directions)
            if supplied_directions
            else await self._latest_user_directions(letter_id)
        )
        if directions_block and directions_block not in (request.points or ""):
            request = request.model_copy(
                update={
                    "points": (
                        f"{request.points}\n\n{directions_block}"
                        if request.points
                        else directions_block
                    )
                }
            )
            trace.append({"stage": "user_direction", "status": "applied"})

        inputs = self._inputs_payload(letter, request)

        # Identity/config fields shared by every terminal DraftRun below. Kept
        # in one place so a new field is added once here, not copied into four
        # hand-written constructor calls that can silently drift apart.
        base_run = dict(
            run_id=run_id,
            letter_id=letter_id,
            draft_type=request.draft_type,
            mode=request.mode,
            letter_category=request.letter_category,
            contract_package=request.contract_package,
            role=role,
            recipient_focus=recipient_focus,
            inputs=inputs,
            started_at=started,
            created_by=self._user_id(current_user),
            idempotency_key=idempotency_key,
            request_hash=request_hash,
            updated_at=started,
        )
        if engine_metadata:
            base_run.update(engine_metadata)

        document_service = DocumentService(self.db)
        context_builder = DraftContextBuilder(
            document_service=document_service,
            conversation_service=ConversationService(self.letter_service),
            db=self.db,
        )

        try:
            context, sources, context_warnings = await context_builder.build(
                letter, request, current_user
            )
            sources = await self._add_governance_comment_context(
                letter_id,
                context,
                sources,
                current_run_id=run_id,
            )
            sources = self._with_source_hashes(sources)
            warnings.extend(context_warnings)
            trace.append({"stage": "context", "status": "success", "source_count": len(sources)})

            # Deterministic input/evidence guardrail scan (H1): user-supplied
            # request text and retrieved source material are untrusted; an
            # injection attempt must be visible to the reviewer, not just
            # neutralised by the prompt guard.
            guardrail_report = self._guardrail_scan(request, sources)
            for finding in guardrail_report.findings:
                warnings.append(f"guardrail {finding.code}: {finding.message}")
            trace.append(
                {
                    "stage": "guardrail_scan",
                    "status": "flagged" if guardrail_report.findings else "success",
                    "verdict": guardrail_report.verdict,
                    "finding_count": len(guardrail_report.findings),
                }
            )

            # The context pack is only persisted on the terminal path actually
            # taken, and its `sources` differ before vs. after generation, so it
            # is built lazily at each terminal below rather than eagerly (and
            # then discarded) on the success path.
            def _context_pack(current_sources: List[SourceEvidence]) -> DraftContextPack:
                return self._build_context_pack(
                    letter=letter,
                    request=request,
                    run_id=run_id,
                    context=context,
                    sources=current_sources,
                    source_warnings=context_warnings,
                )

            incoming_analysis: Optional[IncomingLetterAnalysis] = None
            probing_questions: list[ProbingQuestion] = []
            legal_risk_report: Optional[LegalRiskReport] = None
            if request.draft_type == "reply":
                incoming_analysis = await self.incoming_analyzer.analyze(
                    letter, request, document_service
                )
                trace.append({"stage": "incoming_analysis", "status": "success"})
                # Clause Checking Agent: verify cited clauses against the
                # project's structured clause records before planning.
                try:
                    incoming_analysis = await self.clause_checker.enrich(
                        incoming_analysis,
                        str(getattr(letter, "organization_id", "") or ""),
                        str(getattr(letter, "project_id", "") or ""),
                    )
                    trace.append(
                        {
                            "stage": "clause_check",
                            "status": "success",
                            "verified": sum(
                                1
                                for ev in incoming_analysis.cited_clause_evaluations
                                if ev.exists_in_contract is not None
                            ),
                        }
                    )
                except Exception as exc:
                    warnings.append(f"Clause check skipped: {exc}")
                    trace.append({"stage": "clause_check", "status": "failed"})
                missing_inputs = [
                    key for key, present in context.threshold_inputs.items() if not present
                ]
                probing_questions = UserDirectionAgent.build_questions(
                    incoming_analysis, request, missing_inputs
                )
                if probing_questions:
                    trace.append(
                        {"stage": "user_direction_questions", "count": len(probing_questions)}
                    )

            planning_sheet, reply_matrix, source_summary, deterministic_plan = (
                self.planning_builder.build(
                    letter,
                    request,
                    role,
                    context,
                    sources,
                    incoming_analysis,
                )
            )
            threshold_report = self._merge_reports(
                self.validator.threshold_findings(context),
                self.input_validator.validate_request(letter, request, context, sources),
            )
            if guardrail_report.verdict != "pass":
                threshold_report = self._merge_reports(
                    threshold_report,
                    ValidationReport(
                        blocking=True,
                        findings=[
                            ValidationFinding(
                                level="error",
                                code="critical_prompt_injection",
                                message=(
                                    "Drafting stopped because critical prompt-injection "
                                    "or instruction-manipulation content was detected."
                                ),
                            )
                        ],
                    ),
                )
            if threshold_report.blocking:
                context_pack = _context_pack(sources)
                run = DraftRun(
                    **base_run,
                    status="blocked",
                    incoming_analysis=incoming_analysis,
                    probing_questions=probing_questions,
                    planning_sheet=planning_sheet,
                    reply_matrix=reply_matrix,
                    context_bundle=context,
                    sources=sources,
                    context_pack_id=context_pack.context_pack_id,
                    plan=deterministic_plan,
                    draft_artifact=self._blocked_artifact(threshold_report),
                    source_integrity_summary=source_summary,
                    validation_report=threshold_report,
                    guardrail_report=guardrail_report,
                    evidence_ledger=self._evidence_ledger_entries(run_id, sources),
                    warnings=warnings,
                    trace=trace,
                    completed_at=datetime.now(timezone.utc),
                )
                return await self._create_and_record(run, current_user, context_pack=context_pack)

            required_question_ids = {
                question.question_id for question in probing_questions if question.required
            }
            answered_question_ids = {
                answer.question_id
                for answer in supplied_directions
                if answer.question_id and answer.question_id != "free_text"
            }
            if required_question_ids - answered_question_ids:
                context_pack = _context_pack(sources)
                run = DraftRun(
                    **base_run,
                    status="awaiting_user_direction",
                    execution_status="awaiting_user_direction",
                    next_action="answer_questions",
                    incoming_analysis=incoming_analysis,
                    probing_questions=probing_questions,
                    user_directions=supplied_directions,
                    planning_sheet=planning_sheet,
                    reply_matrix=reply_matrix,
                    context_bundle=context,
                    sources=sources,
                    context_pack_id=context_pack.context_pack_id,
                    plan=deterministic_plan,
                    source_integrity_summary=source_summary,
                    validation_report=threshold_report,
                    guardrail_report=guardrail_report,
                    evidence_ledger=self._evidence_ledger_entries(run_id, sources),
                    warnings=warnings,
                    trace=trace
                    + [
                        {
                            "stage": "user_direction_gate",
                            "status": "interrupted",
                            "missing_question_ids": sorted(
                                required_question_ids - answered_question_ids
                            ),
                        }
                    ],
                )
                return await self._create_and_record(run, current_user, context_pack=context_pack)

            plan = await self._resolve_strategy_plan(letter_id, letter, request)
            if request.mode in _AI_DRAFT_MODES and not plan:
                validation = ValidationReport(
                    blocking=True,
                    findings=[
                        ValidationFinding(
                            level="error",
                            code="missing_strategy_plan",
                            message="Generate and save a strategic plan before AI drafting.",
                        )
                    ],
                )
                context_pack = _context_pack(sources)
                run = DraftRun(
                    **base_run,
                    status="blocked",
                    incoming_analysis=incoming_analysis,
                    probing_questions=probing_questions,
                    planning_sheet=planning_sheet,
                    reply_matrix=reply_matrix,
                    context_bundle=context,
                    sources=sources,
                    context_pack_id=context_pack.context_pack_id,
                    plan=deterministic_plan,
                    draft_artifact=self._blocked_artifact(validation),
                    source_integrity_summary=source_summary,
                    validation_report=validation,
                    guardrail_report=guardrail_report,
                    evidence_ledger=self._evidence_ledger_entries(run_id, sources),
                    warnings=warnings,
                    trace=trace + [{"stage": "strategy", "status": "blocked"}],
                    completed_at=datetime.now(timezone.utc),
                )
                return await self._create_and_record(run, current_user, context_pack=context_pack)
            if request.mode in {"background", "strategy"} and not plan:
                # Constructed lazily: StrategyPlanner builds an AsyncOpenAI
                # client, and only this branch uses it — a plain draft/review
                # run must not pay to build a client it never calls.
                planner = StrategyPlanner(self.prompt_registry)
                plan, prompt_version, plan_warnings = await planner.generate(
                    letter, role, recipient_focus, context, sources, inputs
                )
                plan = plan or deterministic_plan
                warnings.extend(plan_warnings)
                trace.append(
                    {
                        "stage": "strategy",
                        "status": "success",
                        "prompt_version": prompt_version,
                    }
                )

            artifact: Optional[DraftArtifact] = None
            validation = ValidationReport(blocking=False, findings=[])
            cyclic_trace: List[CyclicIterationTrace] = []
            assertion_support: List[DraftAssertionSupport] = []
            confidence_scores: Optional[DraftConfidenceScores] = None
            status = "completed"

            if request.mode in _AI_DRAFT_MODES:
                generator = DraftGenerator(self.prompt_registry)
                artifact, draft_warnings = await generator.generate(
                    letter,
                    role,
                    recipient_focus,
                    context,
                    sources,
                    inputs,
                    plan=plan,
                    finalized=request.finalized,
                )
                warnings.extend(draft_warnings)
                if not stop_after_generation:
                    (
                        artifact,
                        sources,
                        validation,
                        cyclic_trace,
                        assertion_support,
                        confidence_scores,
                        cyclic_warnings,
                    ) = await self._run_cyclic_draft(
                        letter=letter,
                        request=request,
                        current_user=current_user,
                        role=role,
                        recipient_focus=recipient_focus,
                        context=context,
                        sources=sources,
                        inputs=inputs,
                        plan=plan,
                        generator=generator,
                        initial_artifact=artifact,
                    )
                    warnings.extend(cyclic_warnings)
                    if validation.blocking:
                        status = "needs_attention"
                # Fail-visible: the deterministic template fallback (LLM outage
                # or offline mode) is a degraded output that must reach a human
                # as such — never a "completed" run with a buried warning.
                llm_degraded = any(w.startswith("draft_llm:") for w in warnings)
                if llm_degraded:
                    status = "needs_attention"
                trace.append(
                    {
                        "stage": "draft",
                        "status": "degraded" if llm_degraded else "success",
                        "prompt_version": artifact.prompt_version,
                    }
                )
                if not stop_after_generation:
                    trace.append(
                        {
                            "stage": "cyclic_validation",
                            "status": "success",
                            "blocking": validation.blocking,
                            "finding_count": len(validation.findings),
                            "iteration_count": len(cyclic_trace),
                        }
                    )
                # Legal / Contractual Risk Review Agent: flags admissions,
                # waivers, entitlement creation and stance reversals vs the
                # previous position. Flags only — never blocks the run.
                if not stop_after_generation and artifact and artifact.draft_letter:
                    previous_positions = [
                        source.text
                        for source in sources
                        if source.text and (source.metadata or {}).get("previous_position")
                    ]
                    legal_risk_report = self.legal_risk_reviewer.review(
                        artifact.draft_letter, previous_positions
                    )
                    trace.append(
                        {
                            "stage": "legal_risk_review",
                            "status": "success",
                            "flag_count": len(legal_risk_report.flags),
                            "human_review_required": legal_risk_report.human_review_required,
                        }
                    )
            elif request.mode == "background":
                artifact = DraftArtifact(
                    draft_letter="",
                    source_integrity_notes="Background context generated; no draft requested.",
                    raw_model_output="",
                )

            context_pack = _context_pack(sources)

            run = DraftRun(
                **base_run,
                status=status,
                incoming_analysis=incoming_analysis,
                probing_questions=probing_questions,
                user_directions=supplied_directions,
                planning_sheet=planning_sheet,
                reply_matrix=reply_matrix,
                context_bundle=context,
                sources=sources,
                context_pack_id=context_pack.context_pack_id,
                plan=plan,
                draft_artifact=artifact,
                source_integrity_summary=source_summary,
                validation_report=validation,
                legal_risk_report=legal_risk_report,
                guardrail_report=guardrail_report,
                evidence_ledger=self._evidence_ledger_entries(run_id, sources),
                cyclic_trace=cyclic_trace,
                assertion_support=assertion_support,
                confidence_scores=confidence_scores,
                iteration_count=len(cyclic_trace),
                warnings=warnings,
                trace=trace,
                completed_at=datetime.now(timezone.utc),
            )
            return await self._create_and_record(run, current_user, context_pack=context_pack)
        except Exception as exc:
            warnings.append(str(exc))
            run = DraftRun(
                **base_run,
                status="failed",
                warnings=warnings,
                trace=trace + [{"stage": "failed", "status": "error", "message": str(exc)}],
                completed_at=datetime.now(timezone.utc),
            )
            stored = await self._create_and_record(run, current_user)
            await self.repository.append_event(
                letter_id,
                stored.run_id,
                "failed",
                actor_user_id=self._user_id(current_user),
                status=stored.status,
                detail=str(exc),
            )
            return stored

    async def get_run(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.request.view",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        return run

    async def get_audit(self, letter_id: str, run_id: str, current_user: Any) -> DraftAuditResponse:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.audit.view",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        return DraftAuditResponse(
            letter_id=str(letter_id),
            run_id=str(run_id),
            events=await self.repository.list_events(letter_id, run_id),
        )

    async def get_context_pack(
        self,
        letter_id: str,
        run_id: str,
        current_user: Any,
    ) -> DraftContextPack:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.request.view",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        pack = await self.repository.get_context_pack(letter_id, run_id)
        if not pack:
            raise HTTPException(status_code=404, detail="Draft context pack not found")
        return pack

    async def get_source_ledger(
        self,
        letter_id: str,
        run_id: str,
        current_user: Any,
    ) -> SourceLedgerResponse:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.request.view",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        sources = self._with_source_hashes(run.sources)
        return SourceLedgerResponse(
            letter_id=str(letter_id),
            run_id=str(run_id),
            source_count=len(sources),
            sources=sources,
        )

    async def get_governance(
        self,
        letter_id: str,
        run_id: str,
        current_user: Any,
    ) -> DraftGovernanceResponse:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.request.view",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        return DraftGovernanceResponse(
            letter_id=str(letter_id),
            run_id=str(run_id),
            assignments=await self.repository.list_assignments(letter_id, run_id),
            comments=await self.repository.list_comments(letter_id, run_id),
            events=await self.repository.list_events(letter_id, run_id),
        )

    async def get_quality_dashboard(
        self,
        current_user: Any,
        *,
        organization_id: Optional[str] = None,
        project_id: Optional[str] = None,
        window_days: int = 30,
    ) -> DraftQualityDashboardResponse:
        await self.policy_service.authorize(
            current_user,
            "drafting.audit.view",
            resource_type="drafting_dashboard",
            organization_id=organization_id,
            project_id=project_id,
        )
        now = datetime.now(timezone.utc)
        window_days = max(1, min(int(window_days or 30), 365))
        window_start = now - timedelta(days=window_days)
        letter_scope = build_scope_query(
            current_user,
            organization_id=organization_id,
            project_id=project_id,
        )
        letter_docs = await self.db.letters.find(
            letter_scope,
            {"_id": 1},
        ).to_list(length=25000)
        letter_ids = [str(doc.get("_id")) for doc in letter_docs if doc.get("_id")]
        if not letter_ids:
            return DraftQualityDashboardResponse(window_days=window_days)

        run_query: Dict[str, Any] = {
            "letter_id": {"$in": letter_ids},
            "$or": [
                {"started_at": {"$gte": window_start}},
                {"completed_at": {"$gte": window_start}},
                {"approved_at": {"$gte": window_start}},
                {"issued_at": {"$gte": window_start}},
            ],
        }
        run_docs = await self.db.letter_draft_runs.find(run_query).to_list(length=25000)
        runs = [DraftRun(**doc) for doc in run_docs]
        run_ids = [run.run_id for run in runs]
        if not runs:
            return DraftQualityDashboardResponse(window_days=window_days)

        event_docs = await self.db.letter_draft_events.find(
            {
                "letter_id": {"$in": letter_ids},
                "run_id": {"$in": run_ids},
                "created_at": {"$gte": window_start},
            }
        ).to_list(length=50000)
        assignment_docs = await self.db.letter_draft_assignments.find(
            {
                "letter_id": {"$in": letter_ids},
                "run_id": {"$in": run_ids},
                "status": "assigned",
            }
        ).to_list(length=25000)
        issued_docs = await self.db.issued_letters.find(
            {"letter_id": {"$in": letter_ids}, "run_id": {"$in": run_ids}}
        ).to_list(length=25000)

        returned_run_ids = {
            str(event.get("run_id"))
            for event in event_docs
            if event.get("event_type") == "returned_for_correction"
        }
        approved_run_ids = {
            str(event.get("run_id"))
            for event in event_docs
            if event.get("event_type") == "approved"
        }
        issued_artifact_map = {
            str(doc.get("run_id")): bool(doc.get("docx_file_object_id") and doc.get("pdf_file_object_id"))
            for doc in issued_docs
        }

        status_counts: Dict[str, int] = {}
        trend_map: Dict[str, Dict[str, float]] = {}
        durations: List[float] = []
        iterations: List[float] = []
        confidences: List[float] = []
        source_counts: List[int] = []
        unsupported_runs = 0
        blocking_runs = 0
        source_integrity_ok = 0
        source_integrity_total = 0
        active_runs = 0
        approved_runs = 0
        exported_runs = 0
        issued_runs = 0
        recent_risks: List[DraftQualityRiskItem] = []
        bottleneck_ages: Dict[str, List[float]] = {}

        for run in runs:
            status_counts[run.status] = status_counts.get(run.status, 0) + 1
            if run.status in _ACTIVE_STATUSES:
                active_runs += 1
            if run.status in _FINALIZED_STATUSES:
                approved_runs += 1
            if run.status in _EXPORTED_ISSUED_STATUSES:
                exported_runs += 1
            if run.status == "issued":
                issued_runs += 1

            started_at = self._coerce_datetime(run.started_at)
            completed_at = self._coerce_datetime(
                run.issued_at or run.exported_at or run.approved_at or run.completed_at
            )
            if started_at and completed_at and completed_at >= started_at:
                durations.append((completed_at - started_at).total_seconds() / 3600)
            if run.iteration_count is not None:
                iterations.append(float(run.iteration_count))
            if run.confidence_scores:
                confidences.append(float(run.confidence_scores.overall))
            source_count = len(run.sources or [])
            source_counts.append(source_count)
            if run.mode in _AI_DRAFT_MODES:
                source_integrity_total += 1
                if source_count > 0 and all(source.source_hash for source in run.sources):
                    source_integrity_ok += 1

            finding_codes = {
                finding.code
                for finding in (run.validation_report.findings if run.validation_report else [])
            }
            unsupported = any("unsupported" in code for code in finding_codes)
            if unsupported:
                unsupported_runs += 1
            if run.validation_report and run.validation_report.blocking:
                blocking_runs += 1

            trend_key = (started_at or completed_at or now).date().isoformat()
            trend = trend_map.setdefault(
                trend_key,
                {
                    "runs": 0,
                    "approved": 0,
                    "blocking": 0,
                    "unsupported": 0,
                    "confidence_sum": 0.0,
                    "confidence_count": 0,
                },
            )
            trend["runs"] += 1
            if run.status in _FINALIZED_STATUSES:
                trend["approved"] += 1
            if run.validation_report and run.validation_report.blocking:
                trend["blocking"] += 1
            if unsupported:
                trend["unsupported"] += 1
            if run.confidence_scores:
                trend["confidence_sum"] += float(run.confidence_scores.overall)
                trend["confidence_count"] += 1

            if run.status in {"blocked", "needs_attention"}:
                recent_risks.append(
                    DraftQualityRiskItem(
                        run_id=run.run_id,
                        letter_id=run.letter_id,
                        status=run.status,
                        risk="blocking_validation" if run.validation_report.blocking else "needs_attention",
                        detail=self._risk_detail(run),
                        created_at=started_at,
                    )
                )
            if source_count == 0 and run.mode in _AI_DRAFT_MODES:
                recent_risks.append(
                    DraftQualityRiskItem(
                        run_id=run.run_id,
                        letter_id=run.letter_id,
                        status=run.status,
                        risk="missing_sources",
                        detail="Draft/review run has no source ledger entries.",
                        created_at=started_at,
                    )
                )

            age_hours = ((now - (started_at or now)).total_seconds() / 3600)
            if run.status in _ACTIVE_STATUSES:
                bottleneck_ages.setdefault(run.status, []).append(age_hours)

        overdue_review_count = 0
        for assignment in assignment_docs:
            due_at = self._coerce_datetime(assignment.get("due_at"))
            if due_at and due_at < now:
                overdue_review_count += 1

        total_runs = len(runs)
        unsupported_claim_rate = unsupported_runs / total_runs if total_runs else 0.0
        blocking_validation_rate = blocking_runs / total_runs if total_runs else 0.0
        source_integrity_rate = (
            source_integrity_ok / source_integrity_total
            if source_integrity_total
            else 1.0
        )
        review_return_rate = len(returned_run_ids) / total_runs if total_runs else 0.0
        first_review_approval_rate = (
            len(approved_run_ids - returned_run_ids) / len(approved_run_ids)
            if approved_run_ids
            else 0.0
        )
        issue_artifact_compliance_rate = (
            sum(1 for ok in issued_artifact_map.values() if ok) / len(issued_artifact_map)
            if issued_artifact_map
            else (1.0 if issued_runs == 0 else 0.0)
        )
        average_cycle_hours = self._average(durations)
        average_iterations = self._average(iterations)
        average_confidence = self._average(confidences)
        average_sources = self._average([float(count) for count in source_counts])

        response = DraftQualityDashboardResponse(
            window_days=window_days,
            total_runs=total_runs,
            active_runs=active_runs,
            approved_runs=approved_runs,
            exported_runs=exported_runs,
            issued_runs=issued_runs,
            average_cycle_hours=round(average_cycle_hours, 2),
            average_iterations=round(average_iterations, 2),
            average_confidence=round(average_confidence, 4),
            first_review_approval_rate=round(first_review_approval_rate, 4),
            unsupported_claim_rate=round(unsupported_claim_rate, 4),
            blocking_validation_rate=round(blocking_validation_rate, 4),
            source_integrity_rate=round(source_integrity_rate, 4),
            average_sources_per_run=round(average_sources, 2),
            review_return_rate=round(review_return_rate, 4),
            issue_artifact_compliance_rate=round(issue_artifact_compliance_rate, 4),
            overdue_review_count=overdue_review_count,
            status_breakdown=[
                DraftMetricBreakdownItem(
                    label=status,
                    count=count,
                    percentage=round(count / total_runs, 4) if total_runs else 0.0,
                )
                for status, count in sorted(status_counts.items())
            ],
            quality_trends=[
                DraftMetricTrendPoint(
                    date=key,
                    runs=int(value["runs"]),
                    approved=int(value["approved"]),
                    blocking=int(value["blocking"]),
                    unsupported_rate=round(value["unsupported"] / value["runs"], 4)
                    if value["runs"]
                    else 0.0,
                    average_confidence=round(
                        value["confidence_sum"] / value["confidence_count"], 4
                    )
                    if value["confidence_count"]
                    else 0.0,
                )
                for key, value in sorted(trend_map.items())[-14:]
            ],
            bottlenecks=[
                DraftMetricBottleneck(
                    stage=stage,
                    count=len(values),
                    average_age_hours=round(self._average(values), 2),
                )
                for stage, values in sorted(bottleneck_ages.items())
            ],
            recent_risks=sorted(
                recent_risks,
                key=lambda item: item.created_at or datetime.min.replace(tzinfo=timezone.utc),
                reverse=True,
            )[:12],
        )
        response.kpis = self._dashboard_kpis(response)
        return response

    async def assign_reviewer(
        self,
        letter_id: str,
        run_id: str,
        request: AssignReviewerRequest,
        current_user: Any,
    ) -> DraftGovernanceResponse:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.request.assign",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.status in {"blocked", "failed", "approved", "exported", "issued"}:
            raise HTTPException(status_code=400, detail=f"Cannot assign reviewer for draft status '{run.status}'")
        if run.mode not in _AI_DRAFT_MODES or not run.draft_artifact:
            raise HTTPException(status_code=400, detail="Only generated draft runs can be assigned for review")
        if run.created_by and str(run.created_by) == str(request.reviewer_user_id):
            raise HTTPException(status_code=400, detail="Reviewer must be different from drafter")
        assignment = DraftReviewAssignment(
            assignment_id=str(uuid.uuid4()),
            letter_id=str(letter_id),
            run_id=str(run_id),
            reviewer_user_id=request.reviewer_user_id,
            assigned_by=self._user_id(current_user),
            due_at=request.due_at,
            note=request.note,
        )
        await self.repository.upsert_assignment(assignment)
        updated = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "assigned_reviewer_id": request.reviewer_user_id,
                "approval_status": "under_review",
            },
        )
        await self.repository.append_event(
            letter_id,
            run_id,
            "reviewer_assigned",
            actor_user_id=self._user_id(current_user),
            status=(updated or run).status,
            payload={"reviewer_user_id": request.reviewer_user_id, "due_at": request.due_at},
        )
        await self._emit_notification(
            NotificationType.APPROVAL_ASSIGNED,
            letter_id,
            current_user,
            include_users=[request.reviewer_user_id],
            title="Draft review assigned",
            message="A letter draft has been assigned for your review.",
        )
        return await self.get_governance(letter_id, run_id, current_user)

    async def add_comment(
        self,
        letter_id: str,
        run_id: str,
        request: DraftCommentRequest,
        current_user: Any,
    ) -> DraftGovernanceResponse:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.perform",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.assigned_reviewer_id and str(run.assigned_reviewer_id) != str(self._user_id(current_user)):
            can_admin = await self.policy_service.has_permission(current_user, "drafting.admin")
            if not can_admin:
                raise HTTPException(status_code=403, detail="Only the assigned reviewer can comment on this draft")
        comment = DraftReviewComment(
            comment_id=str(uuid.uuid4()),
            letter_id=str(letter_id),
            run_id=str(run_id),
            body=request.body,
            visibility=request.visibility,
            created_by=self._user_id(current_user),
        )
        await self.repository.add_comment(comment)
        await self.repository.append_event(
            letter_id,
            run_id,
            "comment_added",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"comment_id": comment.comment_id, "visibility": request.visibility},
        )
        return await self.get_governance(letter_id, run_id, current_user)

    async def return_for_correction(
        self,
        letter_id: str,
        run_id: str,
        request: ReturnForCorrectionRequest,
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.return_for_revision",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.status in _FINALIZED_STATUSES:
            raise HTTPException(status_code=400, detail="Finalized draft cannot be returned for correction")
        if run.assigned_reviewer_id and str(run.assigned_reviewer_id) != str(self._user_id(current_user)):
            can_admin = await self.policy_service.has_permission(current_user, "drafting.admin")
            if not can_admin:
                raise HTTPException(status_code=403, detail="Only the assigned reviewer can return this draft")
        updated = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "status": "needs_attention",
                "approval_status": "returned_for_correction",
                "returned_reason": request.reason,
                "required_changes": list(request.required_changes),
            },
        )
        result = updated or run
        await self.repository.append_event(
            letter_id,
            run_id,
            "returned_for_correction",
            actor_user_id=self._user_id(current_user),
            status=result.status,
            detail=request.reason,
            payload={"required_changes": list(request.required_changes)},
        )
        await self._emit_notification(
            NotificationType.APPROVAL_REJECTED,
            letter_id,
            current_user,
            include_users=[run.created_by] if run.created_by else None,
            title="Draft returned for correction",
            message=request.reason,
        )
        return result

    async def exact_clause_search(
        self,
        request: ExactClauseSearchRequest,
        current_user: Any,
    ) -> SourceLedgerResponse:
        self._assert_scope_allowed(
            current_user,
            request.organization_id,
            request.project_id,
        )
        response = await ContractService().search_contracts(
            ContractSearchRequest(
                query=request.clause_number,
                organization_id=request.organization_id,
                project_id=request.project_id,
                document_id=request.document_id,
                clause_number=request.clause_number,
                exact_phrase=True,
                limit=request.limit,
                top_docs=min(request.limit, 10),
                chunks_per_doc=3,
                summarize=False,
            ),
            current_user,
        )
        sources: List[SourceEvidence] = []
        for idx, chunk in enumerate(response.results, start=1):
            clause_number = chunk.clause_number or request.clause_number
            sources.append(
                SourceEvidence(
                    source_id=(
                        f"clause:{chunk.document_id or chunk.upload_id or 'unknown'}:"
                        f"{clause_number}:{chunk.chunk_index or idx}"
                    ),
                    source_type="contract_clause",
                    allowed_use="clause",
                    organization_id=request.organization_id,
                    project_id=request.project_id,
                    label=f"{clause_number} {chunk.clause_title or ''}".strip(),
                    text=chunk.text,
                    snippet=self._snippet(chunk.text),
                    document_id=str(chunk.document_id or chunk.upload_id or "") or None,
                    clause_number=clause_number,
                    clause_title=chunk.clause_title,
                    page_numbers=[
                        int(p)
                        for p in (chunk.page_numbers or ([chunk.page] if chunk.page else []))
                        if p
                    ],
                    score=chunk.score,
                    metadata={
                        "upload_id": chunk.upload_id,
                        "file_name": chunk.file_name or chunk.source_filename,
                        "section_heading": chunk.section_heading,
                        "toc_path": chunk.toc_path or [],
                    },
                )
            )
        sources = self._with_source_hashes(sources)
        return SourceLedgerResponse(
            letter_id="",
            run_id="exact-clause",
            source_count=len(sources),
            sources=sources,
        )

    async def exact_reference_search(
        self,
        request: ExactReferenceSearchRequest,
        current_user: Any,
    ) -> SourceLedgerResponse:
        self._assert_scope_allowed(
            current_user,
            request.organization_id,
            request.project_id,
        )
        raw_reference = request.reference.strip()
        normalized = normalize_letter_code(raw_reference)
        exact_regex = f"^{raw_reference}$"
        doc_query = {
            "organization_id": request.organization_id,
            "project_id": request.project_id,
            "$or": [
                {"letterNo": {"$regex": exact_regex, "$options": "i"}},
                {"letter_no": {"$regex": exact_regex, "$options": "i"}},
                {"letterNoNormalized": normalized},
                {"references.letterNo": {"$regex": exact_regex, "$options": "i"}},
                {"references.letter_no": {"$regex": exact_regex, "$options": "i"}},
            ],
        }
        letter_query = {
            "organization_id": request.organization_id,
            "project_id": request.project_id,
            "$or": [
                {"letter_no": {"$regex": exact_regex, "$options": "i"}},
                {"letterNo": {"$regex": exact_regex, "$options": "i"}},
                {"reference": {"$regex": exact_regex, "$options": "i"}},
                {"previous_letter_no": {"$regex": exact_regex, "$options": "i"}},
            ],
        }

        doc_cursor = self.db.documents.find(doc_query).limit(request.limit)
        letter_cursor = self.db.letters.find(letter_query).limit(request.limit)
        docs = [doc async for doc in doc_cursor]
        letters = [letter async for letter in letter_cursor]
        sources: List[SourceEvidence] = []
        for doc in docs:
            doc_id = str(doc.get("_id") or doc.get("document_id") or "")
            letter_no = doc.get("letterNo") or doc.get("letter_no") or raw_reference
            text = doc.get("summary") or doc.get("full_text") or doc.get("ocrText") or doc.get("subject")
            sources.append(
                SourceEvidence(
                    source_id=f"document:{doc_id}",
                    source_type="context_document",
                    allowed_use="fact",
                    organization_id=request.organization_id,
                    project_id=request.project_id,
                    label=str(doc.get("subject") or letter_no or "Referenced document"),
                    text=self._snippet(text, width=1200),
                    snippet=self._snippet(text),
                    document_id=doc_id or None,
                    metadata={
                        "letter_no": letter_no,
                        "uploadType": doc.get("uploadType"),
                        "date": doc.get("date") or doc.get("createdAt"),
                    },
                )
            )
        for letter in letters:
            letter_id = str(letter.get("_id") or letter.get("id") or "")
            letter_no = letter.get("letter_no") or letter.get("letterNo") or raw_reference
            text = letter.get("content") or letter.get("draft_output") or letter.get("subject")
            sources.append(
                SourceEvidence(
                    source_id=f"letter:{letter_id}",
                    source_type="prior_correspondence",
                    allowed_use="history_only",
                    organization_id=request.organization_id,
                    project_id=request.project_id,
                    label=str(letter.get("subject") or letter_no or "Referenced letter"),
                    text=self._snippet(text, width=1200),
                    snippet=self._snippet(text),
                    letter_id=letter_id or None,
                    metadata={
                        "letter_no": letter_no,
                        "status": letter.get("status"),
                        "recipient": letter.get("recipient"),
                    },
                )
            )

        deduped: Dict[str, SourceEvidence] = {}
        for source in self._with_source_hashes(sources):
            deduped[source.source_id] = source
        final_sources = list(deduped.values())[: request.limit]
        return SourceLedgerResponse(
            letter_id="",
            run_id="exact-reference",
            source_count=len(final_sources),
            sources=final_sources,
        )

    async def latest_run(
        self,
        letter_id: str,
        mode: Optional[DraftMode],
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "read",
            drafting_permission="drafting.request.view",
        )
        run = await self.repository.latest(letter_id, mode)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        return run

    async def start_session(
        self,
        request: DraftingStartRequest,
        current_user: Any,
    ) -> DraftingStartResponse:
        letter_id = request.incoming_letter_id or request.incoming_document_id
        if not letter_id:
            raise HTTPException(
                status_code=400,
                detail="Start requires an existing incoming_letter_id or incoming_document_id in this backend version.",
            )
        next_step = "analyze-incoming" if request.draft_type == "reply" else "prepare-plan"
        return DraftingStartResponse(
            session_id=str(uuid.uuid4()),
            letter_id=str(letter_id),
            next_step=next_step,
        )

    async def analyze_incoming(
        self,
        letter_id: str,
        request: DraftRunCreateRequest,
        current_user: Any,
    ) -> DraftRun:
        payload = request.model_copy(update={"mode": "background", "draft_type": "reply"})
        return await self.create_run(letter_id, payload, current_user)

    async def confirm_analysis(
        self,
        letter_id: str,
        run_id: str,
        request: ConfirmAnalysisRequest,
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.edit",
        )
        confirmed = request.analysis.model_copy(
            update={
                "confirmed_by_user": True,
                "confirmed_at": datetime.now(timezone.utc),
            }
        )
        run = await self.repository.update_fields(
            letter_id,
            run_id,
            {"incoming_analysis": confirmed},
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "analysis_confirmed",
            actor_user_id=self._user_id(current_user),
            status=run.status,
        )
        return run

    async def provide_user_direction(
        self,
        letter_id: str,
        run_id: str,
        request: UserDirectionRequest,
        current_user: Any,
    ) -> DraftRun:
        """User Direction Agent: record the drafter's answers / line of action.

        Answers are stored on the run and automatically merged into the inputs
        of subsequent strategy/draft runs for this letter.
        """
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.edit",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if existing.status != "awaiting_user_direction":
            raise HTTPException(
                status_code=409,
                detail="This draft run is not waiting for user direction",
            )
        supplied_by_id = {
            answer.question_id: answer
            for answer in request.answers
            if answer.question_id and answer.question_id != "free_text"
        }
        for question in existing.probing_questions:
            answer = supplied_by_id.get(question.question_id)
            if answer and answer.question_version not in (None, question.question_version):
                raise HTTPException(
                    status_code=409,
                    detail=f"Question '{question.question_id}' changed; refresh before answering",
                )
        answers = list(existing.user_directions or [])
        answers.extend(request.answers)
        if request.directions and request.directions.strip():
            answers.append(
                UserDirectionAnswer(question_id="free_text", answer=request.directions.strip())
            )
        if not answers:
            raise HTTPException(status_code=422, detail="No direction provided")
        answered_ids = {
            answer.question_id
            for answer in answers
            if answer.question_id and answer.question_id != "free_text"
        }
        missing = [
            question.question_id
            for question in existing.probing_questions
            if question.required and question.question_id not in answered_ids
        ]
        if missing:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": "All required probing questions must be answered",
                    "missing_question_ids": missing,
                },
            )
        run = await self.repository.update_fields(
            letter_id,
            run_id,
            {"user_directions": answers},
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "user_direction_provided",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"answer_count": len(answers)},
        )
        payload = {
            key: value
            for key, value in dict(existing.inputs or {}).items()
            if key in DraftRunCreateRequest.model_fields
        }
        payload.update(
            {
                "mode": existing.mode,
                "draft_type": existing.draft_type,
                "letter_category": existing.letter_category,
                "user_direction_answers": answers,
                "user_direction": request.directions,
            }
        )
        resumed = await self.create_run(
            letter_id,
            DraftRunCreateRequest(**payload),
            current_user,
            engine_metadata={
                "engine": existing.engine,
                "engine_version": f"{existing.engine_version}-direction-resume",
                "parent_run_id": existing.run_id,
            },
        )
        await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "status": "completed",
                "execution_status": "completed",
                "next_action": "none",
                "resumed_at": datetime.now(timezone.utc),
            },
        )
        return resumed

    async def lock_paragraphs(
        self,
        letter_id: str,
        run_id: str,
        request: LockParagraphsRequest,
        current_user: Any,
    ) -> DraftRun:
        """Lock human-approved paragraphs so redrafts cannot change them."""
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.edit",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        locked = [p.strip() for p in request.locked_paragraphs if p and p.strip()]
        # Warn (without rejecting) when a lock does not match the current
        # draft — the human may have lightly edited the paragraph first.
        draft_text = existing.draft_artifact.draft_letter if existing.draft_artifact else ""
        unmatched = verify_locked_paragraphs(draft_text, locked) if draft_text else []
        run = await self.repository.update_fields(
            letter_id, run_id, {"locked_paragraphs": locked}
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "paragraphs_locked",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"locked_count": len(locked), "unmatched_in_draft": len(unmatched)},
        )
        return run

    async def freeze_sections(
        self,
        letter_id: str,
        run_id: str,
        request: FreezeSectionsRequest,
        current_user: Any,
    ) -> DraftRun:
        """Capture selected sections from the stored draft as immutable text."""

        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.edit",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if existing.status in _FINALIZED_STATUSES:
            raise HTTPException(status_code=409, detail="Finalized drafts cannot change frozen sections")
        draft_text = existing.draft_artifact.draft_letter if existing.draft_artifact else ""
        if not draft_text:
            raise HTTPException(status_code=400, detail="Run does not contain a draft artifact")
        current_hash = content_hash(draft_text)
        if request.expected_draft_hash and request.expected_draft_hash != current_hash:
            raise HTTPException(status_code=409, detail="Draft changed; refresh section selection and retry")
        try:
            frozen = capture_frozen_sections(
                draft_text,
                request.section_indices,
                frozen_by=self._user_id(current_user),
            )
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        run = await self.repository.update_fields(
            letter_id,
            run_id,
            {"frozen_sections": frozen},
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "sections_frozen",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={
                "section_indices": [item.section_index for item in frozen],
                "draft_hash": current_hash,
            },
        )
        return run

    async def _latest_user_directions(self, letter_id: str) -> Optional[str]:
        """Most recent run's user directions, formatted for drafting inputs."""
        try:
            latest = await self.repository.latest(letter_id, None)
        except Exception:
            return None
        if not latest or not latest.user_directions:
            return None
        return UserDirectionAgent.format_directions(latest.user_directions) or None

    async def prepare_plan(
        self,
        letter_id: str,
        request: DraftRunCreateRequest,
        current_user: Any,
    ) -> DraftRun:
        payload = request.model_copy(update={"mode": "strategy"})
        run = await self.create_run(letter_id, payload, current_user)
        if run.plan and run.status not in {"failed", "blocked"}:
            version = await self.repository.save_strategy_plan(
                letter_id,
                run,
                self._user_id(current_user),
                status="generated",
            )
            await self.repository.append_event(
                letter_id,
                run.run_id,
                "plan_confirmed",
                actor_user_id=self._user_id(current_user),
                status=run.status,
                detail="Strategic plan generated and saved as latest version.",
                payload={"strategy_version": version, "auto_saved": True},
            )
        return run

    async def confirm_plan(
        self,
        letter_id: str,
        run_id: str,
        request: ConfirmPlanRequest,
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.draft.create",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        fields: Dict[str, Any] = {}
        if request.planning_sheet:
            fields["planning_sheet"] = request.planning_sheet.model_copy(
                update={"user_confirmed": True}
            )
        if request.reply_matrix:
            fields["reply_matrix"] = request.reply_matrix
        if request.planning_sheet:
            fields["plan"] = self.planning_builder._plan_text(
                fields["planning_sheet"],
                request.reply_matrix or existing.reply_matrix,
            )
        fields["approval_status"] = "plan_confirmed"
        run = await self.repository.update_fields(letter_id, run_id, fields)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.plan:
            version = await self.repository.save_strategy_plan(
                letter_id,
                run,
                self._user_id(current_user),
                status="confirmed",
            )
        else:
            version = None
        await self.repository.append_event(
            letter_id,
            run_id,
            "plan_confirmed",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"strategy_version": version},
        )
        return run

    async def generate_draft(
        self,
        letter_id: str,
        request: DraftRunCreateRequest,
        current_user: Any,
    ) -> DraftRun:
        payload = request.model_copy(update={"mode": "draft"})
        return await self.create_run(letter_id, payload, current_user)

    async def revise_run(
        self,
        letter_id: str,
        run_id: str,
        request: ReviseDraftRequest,
        current_user: Any,
    ) -> DraftRun:
        source = await self.get_run(letter_id, run_id, current_user)
        # Locked text: explicit list on the request wins ([] clears), otherwise
        # locks carry forward from the run being revised.
        locked_paragraphs = (
            request.locked_paragraphs
            if request.locked_paragraphs is not None
            else list(source.locked_paragraphs or [])
        )
        revision_text = self._revision_instruction(request)
        locks_block = locked_instruction(locked_paragraphs)
        if locks_block:
            revision_text = f"{revision_text}\n\n{locks_block}"
        original_inputs = source.inputs or {}
        payload = DraftRunCreateRequest(
            mode="draft",
            draft_type=source.draft_type,
            letter_category=source.letter_category,
            contract_package=source.contract_package,
            role=source.role,
            recipient_focus=source.recipient_focus,
            subject=original_inputs.get("subject"),
            recipient=original_inputs.get("recipient"),
            requirements=revision_text,
            points=original_inputs.get("points"),
            purpose=original_inputs.get("purpose"),
            desired_position=original_inputs.get("desired_position"),
            required_action=original_inputs.get("required_action"),
            background_facts=original_inputs.get("background_facts"),
            trigger_event=original_inputs.get("trigger_event"),
            tone=original_inputs.get("tone") or "firm_contractual",
            timeline_days=original_inputs.get("timeline_days"),
            incoming_document_id=original_inputs.get("incoming_document_id"),
            incoming_letter_id=original_inputs.get("incoming_letter_id"),
            clauses_to_consider=original_inputs.get("clauses_to_consider") or [],
            attachments=original_inputs.get("attachments") or [],
            document_ids=original_inputs.get("document_ids") or [],
            include_letter_codes=original_inputs.get("include_letter_codes") or [],
            exclude_letter_codes=original_inputs.get("exclude_letter_codes") or [],
            plan_override=source.plan,
        )
        revised = await self.create_run(letter_id, payload, current_user)

        # Verify human-locked paragraphs survived verbatim; retry once with a
        # stronger instruction, then surface (never silently repair).
        locked_violations: list[str] = []
        if locked_paragraphs:
            draft_text = revised.draft_artifact.draft_letter if revised.draft_artifact else ""
            locked_violations = verify_locked_paragraphs(draft_text, locked_paragraphs)
            if locked_violations:
                retry_payload = payload.model_copy(
                    update={
                        "requirements": (
                            f"{revision_text}\n\nIMPORTANT: the previous attempt modified "
                            "locked paragraphs. Reproduce every LOCKED paragraph exactly as "
                            "provided, character for character."
                        )
                    }
                )
                retry = await self.create_run(letter_id, retry_payload, current_user)
                retry_text = retry.draft_artifact.draft_letter if retry.draft_artifact else ""
                retry_violations = verify_locked_paragraphs(retry_text, locked_paragraphs)
                if len(retry_violations) < len(locked_violations):
                    revised = retry
                    locked_violations = retry_violations

        update_fields: Dict[str, Any] = {
            "revision_of_run_id": run_id,
            "revision_action": request.revision_action,
            "locked_paragraphs": locked_paragraphs,
        }
        if locked_violations:
            update_fields["warnings"] = list(revised.warnings or []) + [
                "Locked paragraph(s) were modified by the AI and could not be preserved "
                f"after retry ({len(locked_violations)} affected). Restore the locked text "
                "before approval."
            ]
        updated = await self.repository.update_fields(letter_id, revised.run_id, update_fields)
        result = updated or revised
        await self.repository.append_event(
            letter_id,
            result.run_id,
            "revised",
            actor_user_id=self._user_id(current_user),
            status=result.status,
            payload={
                "revision_of_run_id": run_id,
                "revision_action": request.revision_action,
                "locked_paragraph_count": len(locked_paragraphs),
                "locked_violations": len(locked_violations),
            },
        )
        return result

    async def revise_selected_sections(
        self,
        letter_id: str,
        run_id: str,
        request: ReviseSectionsRequest,
        current_user: Any,
    ) -> DraftRun:
        """Rewrite only selected non-frozen sections of a stored draft.

        Frozen and unselected text never enters the model-editable response
        path. The server reconstructs the result from the source draft and then
        blocks approval if any preserved section fails exact verification.
        """

        source = await self.get_run(letter_id, run_id, current_user)
        if source.status in _FINALIZED_STATUSES:
            raise HTTPException(status_code=409, detail="Finalized drafts cannot be revised")
        if not source.draft_artifact or not source.draft_artifact.draft_letter:
            raise HTTPException(status_code=400, detail="Run does not contain a draft artifact")

        source_text = source.draft_artifact.draft_letter
        source_hash = content_hash(source_text)
        if request.expected_draft_hash and request.expected_draft_hash != source_hash:
            raise HTTPException(status_code=409, detail="Draft changed; refresh section selection and retry")

        available = section_map(source_text)
        editable_indices = sorted(set(int(index) for index in request.section_indices))
        unknown = [index for index in editable_indices if index not in available]
        if unknown:
            raise HTTPException(
                status_code=422,
                detail=f"Unknown draft section indices: {unknown}",
            )

        frozen_violations = verify_frozen_sections(source_text, source.frozen_sections)
        if frozen_violations:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "Stored frozen section integrity check failed",
                    "section_indices": frozen_violations,
                },
            )
        frozen_indices = {item.section_index for item in source.frozen_sections}
        conflicts = sorted(frozen_indices.intersection(editable_indices))
        if conflicts:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "Frozen sections cannot be selected for editing",
                    "section_indices": conflicts,
                },
            )

        editor = ScopedSectionEditor()
        replacements, edit_warnings = await editor.edit(
            {index: available[index].content for index in editable_indices},
            request.action,
        )
        anchor_violations: Dict[int, Dict[str, List[str]]] = {}
        for index, replacement in replacements.items():
            changes = protected_anchor_changes(available[index].content, replacement)
            if changes["removed"] or changes["introduced"]:
                anchor_violations[index] = changes
        if anchor_violations:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": "Scoped edit changed protected factual anchors",
                    "sections": anchor_violations,
                },
            )
        try:
            revised_text = apply_section_edits(source_text, replacements)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

        preserved = preserved_sections(source_text, editable_indices)
        preservation_violations = verify_frozen_sections(revised_text, preserved)
        if preservation_violations:
            raise HTTPException(
                status_code=500,
                detail={
                    "message": "Scoped revision integrity check failed",
                    "section_indices": preservation_violations,
                },
            )

        artifact = source.draft_artifact.model_copy(
            update={
                "draft_letter": revised_text,
                "raw_model_output": revised_text,
            }
        )
        validation = self._merge_reports(
            self.validator.critique(
                artifact,
                source.role,
                source.sources,
                finalized=bool((source.inputs or {}).get("finalized")),
            ),
            self.validator.strategy_alignment(artifact, source.plan or ""),
        )
        previous_positions = [
            evidence.text or evidence.snippet or ""
            for evidence in source.sources
            if evidence.source_type in {"prior_correspondence", "graph_thread"}
        ]
        legal_risk_report = self.legal_risk_reviewer.review(
            artifact.draft_letter,
            previous_positions=previous_positions,
        )
        now = datetime.now(timezone.utc)
        warnings = list(source.warnings or []) + list(edit_warnings)
        status = "needs_attention" if validation.blocking or edit_warnings else "completed"
        section_revision = SectionRevisionRecord(
            source_run_id=source.run_id,
            action=request.action,
            editable_section_indices=editable_indices,
            preserved_sections=preserved,
            source_draft_hash=source_hash,
            result_draft_hash=content_hash(revised_text),
            prompt_version=SECTION_EDIT_PROMPT_VERSION,
        )
        revised_inputs = dict(source.inputs or {})
        revised_inputs["section_edit"] = {
            "source_run_id": source.run_id,
            "action": request.action,
            "editable_section_indices": editable_indices,
            "source_draft_hash": source_hash,
            "result_draft_hash": section_revision.result_draft_hash,
            "prompt_version": SECTION_EDIT_PROMPT_VERSION,
        }
        revised = source.model_copy(
            update={
                "id": None,
                "run_id": str(uuid.uuid4()),
                "status": status,
                "inputs": revised_inputs,
                "draft_artifact": artifact,
                "validation_report": validation,
                "legal_risk_report": legal_risk_report,
                "warnings": warnings,
                "revision_of_run_id": source.run_id,
                "revision_action": None,
                "section_revision": section_revision,
                "approvals": [],
                "approval_status": None,
                "approved_by": None,
                "approved_at": None,
                "returned_reason": None,
                "required_changes": [],
                "issued_document_id": None,
                "exported_file_id": None,
                "exported_pdf_file_id": None,
                "exported_docx_file_id": None,
                "exported_by": None,
                "exported_at": None,
                "issued_by": None,
                "issued_at": None,
                "engine": "v2",
                "engine_version": "v2-section-edit",
                "graph_version": None,
                "thread_id": None,
                "idempotency_key": None,
                "request_hash": None,
                "attempt_number": 1,
                "parent_run_id": source.run_id,
                "fallback_of_run_id": None,
                "fallback_reason": None,
                "shadow_of_run_id": None,
                "context_snapshot_id": None,
                "input_snapshot_id": None,
                "input_snapshot_hash": None,
                "context_snapshot_hash": None,
                "execution_status": "completed",
                "next_action": "approve",
                "state_version": 0,
                "last_checkpoint_id": None,
                "queue_job_id": None,
                "lease_owner": None,
                "lease_expires_at": None,
                "cancellation_requested_at": None,
                "cancellation_reason": None,
                "resumed_at": None,
                "trace": list(source.trace or [])
                + [
                    {
                        "stage": "section_revision",
                        "status": "success" if not edit_warnings else "degraded",
                        "action": request.action,
                        "editable_section_indices": editable_indices,
                        "preserved_section_count": len(preserved),
                    }
                ],
                "started_at": now,
                "completed_at": now,
                "updated_at": now,
                "created_by": self._user_id(current_user),
            }
        )
        stored = await self._create_and_record(revised, current_user)
        await self.repository.append_event(
            letter_id,
            stored.run_id,
            "sections_revised",
            actor_user_id=self._user_id(current_user),
            status=stored.status,
            payload={
                "source_run_id": source.run_id,
                "action": request.action,
                "editable_section_indices": editable_indices,
                "preserved_section_count": len(preserved),
                "source_draft_hash": source_hash,
                "result_draft_hash": section_revision.result_draft_hash,
            },
        )
        return stored

    async def validate_run(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.perform",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if not run.draft_artifact:
            raise HTTPException(status_code=400, detail="Run does not contain a draft artifact")

        validation = self._merge_reports(
            self.validator.validate(
                run.draft_artifact,
                run.role,
                run.sources,
                finalized=bool((run.inputs or {}).get("finalized")),
            ),
            self.validator.strategy_alignment(run.draft_artifact, run.plan or ""),
        )
        integrity_violations = self._section_integrity_violations(run)
        if integrity_violations:
            validation = self._merge_reports(
                validation,
                ValidationReport(
                    blocking=True,
                    findings=[
                        ValidationFinding(
                            level="error",
                            code="frozen_section_modified",
                            message="Frozen or unselected draft content no longer matches its approved exact text.",
                            evidence=", ".join(str(index) for index in integrity_violations),
                        )
                    ],
                ),
            )
        status = "needs_attention" if validation.blocking else run.status
        if run.status in {"blocked", "failed"} and not validation.blocking:
            status = "completed"
        updated = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "validation_report": validation,
                "status": status,
                "last_validated_at": datetime.now(timezone.utc),
            },
        )
        result = updated or run
        await self.repository.append_event(
            letter_id,
            run_id,
            "validated",
            actor_user_id=self._user_id(current_user),
            status=result.status,
            payload={
                "blocking": validation.blocking,
                "finding_count": len(validation.findings),
            },
        )
        return result

    async def critique_run(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.perform",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if not run.draft_artifact:
            raise HTTPException(status_code=400, detail="Run does not contain a draft artifact")

        critique = self._merge_reports(
            self.validator.critique(
                run.draft_artifact,
                run.role,
                run.sources,
                finalized=bool((run.inputs or {}).get("finalized")),
            ),
            self.validator.strategy_alignment(run.draft_artifact, run.plan or ""),
        )
        status = "needs_attention" if critique.blocking else run.status
        trace = list(run.trace or [])
        trace.append(
            {
                "stage": "critique",
                "status": "success",
                "blocking": critique.blocking,
                "finding_count": len(critique.findings),
                "completed_at": datetime.now(timezone.utc).isoformat(),
            }
        )
        updated = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "validation_report": critique,
                "status": status,
                "trace": trace,
                "last_validated_at": datetime.now(timezone.utc),
            },
        )
        result = updated or run
        await self.repository.append_event(
            letter_id,
            run_id,
            "critiqued",
            actor_user_id=self._user_id(current_user),
            status=result.status,
            payload={
                "blocking": critique.blocking,
                "finding_count": len(critique.findings),
            },
        )
        return result

    # Approval chain (multi-agent workflow Phase 4):
    # drafter -> reviewer -> final, each stage its own permission + actor.
    _APPROVAL_STAGE_ORDER: tuple = ("drafter", "reviewer", "final")
    _APPROVAL_STAGE_PERMISSION: dict = {
        "drafter": "drafting.draft.submit_for_review",
        "reviewer": "drafting.review.approve",
        "final": "drafting.final.approve",
    }
    _APPROVAL_STAGE_STATUS: dict = {
        "drafter": "drafter_approved",
        "reviewer": "reviewer_approved",
        "final": "approved",
    }

    @classmethod
    def next_approval_stage(cls, approvals: List[ApprovalStep]) -> Optional[ApprovalStage]:
        done = [step.stage for step in approvals or []]
        for stage in cls._APPROVAL_STAGE_ORDER:
            if stage not in done:
                return stage  # type: ignore[return-value]
        return None

    async def approve_stage(
        self,
        letter_id: str,
        run_id: str,
        request: ApproveStageRequest,
        current_user: Any,
    ) -> DraftRun:
        """One step of the drafter -> reviewer -> final approval chain.

        Order is enforced; each stage requires its own permission and a
        different actor (separation of duties), with drafting.admin as the
        escape hatch. Only the final stage locks the draft version and marks
        the run approved.
        """
        stage = request.stage
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission=self._APPROVAL_STAGE_PERMISSION[stage],
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        already_approved = next(
            (step for step in existing.approvals if step.stage == stage),
            None,
        )
        if already_approved:
            return existing
        if existing.status in _FINALIZED_STATUSES:
            raise HTTPException(status_code=400, detail="Draft is already approved or finalized")
        if existing.status in {"blocked", "failed", "needs_attention"}:
            raise HTTPException(status_code=400, detail="Draft must pass validation before approval")
        if (
            existing.legal_risk_report
            and existing.legal_risk_report.human_review_required
            and not existing.legal_risk_report.human_reviewed_at
        ):
            raise HTTPException(
                status_code=409,
                detail="High legal-risk flags require recorded human review before approval",
            )
        if existing.mode not in _AI_DRAFT_MODES or not existing.draft_artifact:
            raise HTTPException(status_code=400, detail="Run does not contain an approvable draft")
        if not (existing.plan or "").strip():
            raise HTTPException(
                status_code=400, detail="Draft cannot be approved without a saved strategic plan"
            )
        if existing.validation_report.blocking:
            raise HTTPException(status_code=400, detail="Draft has blocking validation findings")
        if existing.guardrail_report and existing.guardrail_report.verdict != "pass":
            raise HTTPException(
                status_code=409,
                detail="Critical guardrail findings must be removed before approval",
            )
        integrity_violations = self._section_integrity_violations(existing)
        if integrity_violations:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "Frozen or preserved section integrity check failed",
                    "section_indices": integrity_violations,
                },
            )

        expected = self.next_approval_stage(existing.approvals)
        if expected is None:
            raise HTTPException(status_code=409, detail="Approval chain is already complete")
        if stage != expected:
            raise HTTPException(
                status_code=409,
                detail=f"Approval stages must run in order; next expected stage is '{expected}'",
            )

        actor = self._user_id(current_user)
        is_admin = await self.policy_service.has_permission(current_user, "drafting.admin")
        prior_actors = {step.stage: step.approved_by for step in existing.approvals}

        if stage == "drafter":
            # The drafter's own sign-off: only the run creator (or admin).
            if (
                existing.created_by
                and str(existing.created_by) != str(actor)
                and not is_admin
            ):
                raise HTTPException(
                    status_code=403,
                    detail="Drafter approval must come from the draft's creator",
                )
        else:
            # Separation of duties: reviewer/final must differ from the earlier
            # actors in the chain (admin may override).
            if not is_admin:
                conflicting = {prior_actors.get("drafter")}
                if stage == "final":
                    conflicting.add(prior_actors.get("reviewer"))
                if str(actor) in {str(a) for a in conflicting if a}:
                    raise HTTPException(
                        status_code=403,
                        detail=f"{stage.capitalize()} approval requires a different approver",
                    )
            if (
                stage == "reviewer"
                and existing.assigned_reviewer_id
                and str(existing.assigned_reviewer_id) != str(actor)
                and not is_admin
            ):
                raise HTTPException(
                    status_code=403,
                    detail="Reviewer approval is restricted to the assigned reviewer",
                )

        now = datetime.now(timezone.utc)
        approvals = list(existing.approvals or []) + [
            ApprovalStep(stage=stage, approved_by=actor, approved_at=now, comment=request.comment)
        ]
        fields: Dict[str, Any] = {
            "approvals": approvals,
            "approval_status": self._APPROVAL_STAGE_STATUS[stage],
        }
        approved_version = None
        if stage == "final":
            run, approved_version = await self.repository.finalize_draft_approval(
                letter_id,
                existing,
                approvals,
                actor,
                now,
            )
        else:
            run = await self.repository.update_fields(letter_id, run_id, fields)
            if not run:
                raise HTTPException(status_code=404, detail="Draft run not found")

        await self.repository.append_event(
            letter_id,
            run_id,
            f"{stage}_approved",
            actor_user_id=actor,
            status=run.status,
            payload={"comment": request.comment} if request.comment else None,
        )
        if stage == "final":
            # Legacy audit consumers count "approved" events — keep emitting it.
            await self.repository.append_event(
                letter_id,
                run_id,
                "approved",
                actor_user_id=actor,
                status=run.status,
                payload={"approved_draft_version": approved_version},
            )
            await self._emit_notification(
                NotificationType.DRAFT_APPROVED,
                letter_id,
                current_user,
                include_users=[existing.created_by] if existing.created_by else None,
                title="Draft approved",
                message="A letter draft has completed the approval chain.",
            )
        return run

    async def review_legal_risk(
        self,
        letter_id: str,
        run_id: str,
        comment: Optional[str],
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.perform",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if not run.legal_risk_report or not run.legal_risk_report.human_review_required:
            raise HTTPException(status_code=409, detail="This run has no mandatory legal-risk review")
        report = run.legal_risk_report.model_copy(
            update={
                "human_reviewed_at": datetime.now(timezone.utc),
                "human_reviewed_by": self._user_id(current_user),
                "human_review_comment": comment,
            }
        )
        updated = await self.repository.update_fields(letter_id, run_id, {"legal_risk_report": report})
        if not updated:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "legal_risk_reviewed",
            actor_user_id=self._user_id(current_user),
            status=updated.status,
            detail=comment,
        )
        return updated

    async def approve_run(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        """Legacy single-approve endpoint: advances the next pending stage of
        the drafter -> reviewer -> final chain."""
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        stage = self.next_approval_stage(existing.approvals)
        if stage is None:
            raise HTTPException(status_code=409, detail="Approval chain is already complete")
        return await self.approve_stage(
            letter_id, run_id, ApproveStageRequest(stage=stage), current_user
        )

    async def approve_legacy_workflow_stage(
        self,
        letter_id: str,
        stage: ApprovalStage,
        letter_content: str,
        current_user: Any,
        *,
        run_id: Optional[str] = None,
        expected_draft_hash: Optional[str] = None,
        comment: Optional[str] = None,
    ) -> DraftRun:
        """Bind a legacy status transition to the exact governed run artifact."""

        letter = await self._load_and_authorize(letter_id, current_user, "write")
        selected_run_id = run_id
        if not selected_run_id and stage != "drafter":
            selected_run_id = getattr(letter, "graph_run_id", None)
        run = (
            await self.repository.get(letter_id, selected_run_id)
            if selected_run_id
            else await self.repository.latest(letter_id, mode="draft")
        )
        if not run or not run.draft_artifact:
            raise HTTPException(
                status_code=409,
                detail="A governed draft run is required for this workflow transition",
            )
        if stage != "drafter" and str(getattr(letter, "graph_run_id", "") or "") != run.run_id:
            raise HTTPException(
                status_code=409,
                detail="The letter is not bound to this governed draft run",
            )
        artifact_body = run.draft_artifact.draft_letter
        artifact_hash = hashlib.sha256(artifact_body.encode("utf-8")).hexdigest()
        mutable_hash = hashlib.sha256((letter_content or "").encode("utf-8")).hexdigest()
        if expected_draft_hash and expected_draft_hash != artifact_hash:
            raise HTTPException(
                status_code=409,
                detail="The supplied draft hash does not match the governed artifact",
            )
        if mutable_hash != artifact_hash:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Letter content changed outside the governed draft run. "
                    "Create a governed revision before continuing."
                ),
            )
        if stage != "drafter":
            bound_hash = str(getattr(letter, "governed_draft_hash", "") or "")
            if bound_hash and bound_hash != artifact_hash:
                raise HTTPException(
                    status_code=409,
                    detail="The governed artifact changed after submission",
                )
        approved = await self.approve_stage(
            letter_id,
            run.run_id,
            ApproveStageRequest(stage=stage, comment=comment),
            current_user,
        )
        if stage == "drafter":
            await self.repository.accept_draft(
                letter_id,
                approved,
                self._user_id(current_user),
            )
        return approved

    async def export_run(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.final.view",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if existing.status not in _FINALIZED_STATUSES:
            raise HTTPException(status_code=400, detail="Draft must be approved before export")
        if existing.status in _EXPORTED_ISSUED_STATUSES and existing.exported_docx_file_id and existing.exported_pdf_file_id:
            return existing
        effect_key = f"{run_id}:export"
        prior_effect = await self.repository.get_effect(effect_key)
        if prior_effect and prior_effect.status == "completed":
            return existing
        await self.repository.record_effect(
            DraftExecutionEffect(
                effect_id=str(uuid.uuid4()), effect_key=effect_key, run_id=run_id,
                effect_type="export_artifacts",
                payload_hash=hashlib.sha256((existing.draft_artifact.draft_letter if existing.draft_artifact else "").encode("utf-8")).hexdigest(),
            )
        )
        now = datetime.now(timezone.utc)
        export_files = await self._export_artifacts(letter_id, existing, current_user)
        run = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "status": "exported",
                "exported_file_id": export_files.get("pdf_file_object_id")
                or export_files.get("docx_file_object_id")
                or f"draft-export:{run_id}",
                "exported_docx_file_id": export_files.get("docx_file_object_id"),
                "exported_pdf_file_id": export_files.get("pdf_file_object_id"),
                "exported_by": self._user_id(current_user),
                "exported_at": now,
            },
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.repository.append_event(
            letter_id,
            run_id,
            "exported",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload=export_files,
        )
        await self.repository.complete_effect(effect_key)
        await self.repository.enqueue_outbox(
            DraftOutboxEvent(
                event_id=f"{run_id}:exported", run_id=run_id, event_type="draft_exported",
                payload={"docx_file_object_id": run.exported_docx_file_id, "pdf_file_object_id": run.exported_pdf_file_id},
            )
        )
        return run

    async def issue_run(
        self,
        letter_id: str,
        run_id: str,
        issued_document_id: Optional[str],
        current_user: Any,
    ) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.final.view",
        )
        existing = await self.repository.get(letter_id, run_id)
        if not existing:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if existing.status not in _EXPORTED_ISSUED_STATUSES:
            raise HTTPException(status_code=400, detail="Draft must be exported before issue")
        if existing.status == "issued":
            return existing
        if not existing.exported_docx_file_id or not existing.exported_pdf_file_id:
            raise HTTPException(
                status_code=400,
                detail="Draft must have immutable DOCX and PDF export artifacts before issue",
            )
        effect_key = f"{run_id}:issue:{issued_document_id or ''}"
        prior_effect = await self.repository.get_effect(effect_key)
        if prior_effect and prior_effect.status == "completed":
            return existing
        await self.repository.record_effect(
            DraftExecutionEffect(
                effect_id=str(uuid.uuid4()), effect_key=effect_key, run_id=run_id,
                effect_type="issue_draft",
                payload_hash=hashlib.sha256(str(issued_document_id or "").encode("utf-8")).hexdigest(),
            )
        )
        now = datetime.now(timezone.utc)
        issued_id = issued_document_id or f"issued:{run_id}"
        run = await self.repository.update_fields(
            letter_id,
            run_id,
            {
                "status": "issued",
                "issued_document_id": issued_id,
                "issued_by": self._user_id(current_user),
                "issued_at": now,
            },
        )
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        await self.db.issued_letters.update_one(
            {"letter_id": str(letter_id), "run_id": str(run_id)},
            {
                "$set": {
                    "issued_letter_id": issued_id,
                    "letter_id": str(letter_id),
                    "run_id": str(run_id),
                    "issued_document_id": issued_id,
                    "exported_file_id": existing.exported_file_id,
                    "docx_file_object_id": existing.exported_docx_file_id,
                    "pdf_file_object_id": existing.exported_pdf_file_id,
                    "source_ledger_hash": self._source_ledger_hash(existing.sources),
                    "issued_by": self._user_id(current_user),
                    "issued_at": now,
                    "status": "issued",
                    "updated_at": now,
                }
            },
            upsert=True,
        )
        await self.repository.append_event(
            letter_id,
            run_id,
            "issued",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"issued_document_id": issued_id},
        )
        await self.repository.complete_effect(effect_key)
        await self.repository.enqueue_outbox(
            DraftOutboxEvent(
                event_id=f"{run_id}:issued:{issued_id}", run_id=run_id, event_type="draft_issued",
                payload={"issued_document_id": issued_id},
            )
        )
        await self._emit_notification(
            NotificationType.APPROVAL_COMPLETED,
            letter_id,
            current_user,
            include_users=[existing.created_by] if existing.created_by else None,
            title="Draft issued",
            message="The approved letter draft has been issued.",
        )
        return run

    async def accept_plan(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.request.accept",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.mode not in {"strategy", "draft", "review", "background"} or not run.plan:
            raise HTTPException(status_code=400, detail="Run does not contain an acceptable plan")
        version = await self.repository.mark_accepted_plan(letter_id, run, self._user_id(current_user))
        await self.repository.append_event(
            letter_id,
            run_id,
            "plan_accepted",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"strategy_version": version},
        )
        return run

    async def accept_draft(self, letter_id: str, run_id: str, current_user: Any) -> DraftRun:
        await self._load_and_authorize(
            letter_id,
            current_user,
            "write",
            drafting_permission="drafting.review.approve",
        )
        run = await self.repository.get(letter_id, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Draft run not found")
        if run.mode not in _AI_DRAFT_MODES or not run.draft_artifact:
            raise HTTPException(status_code=400, detail="Run does not contain an acceptable draft")
        if not (run.plan or "").strip():
            raise HTTPException(status_code=400, detail="Draft cannot be accepted without a saved strategic plan")
        if run.created_by and str(run.created_by) == str(self._user_id(current_user)):
            can_self_approve = await self.policy_service.has_permission(current_user, "drafting.admin")
            if not can_self_approve:
                raise HTTPException(status_code=403, detail="Drafters cannot approve their own draft")
        if run.validation_report.blocking:
            raise HTTPException(status_code=400, detail="Draft has blocking validation findings")
        if run.guardrail_report and run.guardrail_report.verdict != "pass":
            raise HTTPException(status_code=409, detail="Draft has critical guardrail findings")
        if (
            run.legal_risk_report
            and run.legal_risk_report.human_review_required
            and not run.legal_risk_report.human_reviewed_at
        ):
            raise HTTPException(
                status_code=409,
                detail="High legal-risk flags require recorded human review before acceptance",
            )
        integrity_violations = self._section_integrity_violations(run)
        if integrity_violations:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "Frozen or preserved section integrity check failed",
                    "section_indices": integrity_violations,
                },
            )
        version = await self.repository.accept_draft(letter_id, run, self._user_id(current_user))
        await self.repository.append_event(
            letter_id,
            run_id,
            "draft_accepted",
            actor_user_id=self._user_id(current_user),
            status=run.status,
            payload={"draft_version": version},
        )
        return run

    @staticmethod
    def _average(values: List[float]) -> float:
        return sum(values) / len(values) if values else 0.0

    @staticmethod
    def _coerce_datetime(value: Any) -> Optional[datetime]:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        if isinstance(value, str):
            try:
                parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
            except Exception:
                return None
        return None

    @staticmethod
    def _risk_detail(run: DraftRun) -> str:
        findings = run.validation_report.findings if run.validation_report else []
        if findings:
            return findings[0].message
        if run.returned_reason:
            return run.returned_reason
        return "Run requires reviewer attention."

    @staticmethod
    def _dashboard_kpis(response: DraftQualityDashboardResponse) -> List[DraftMetricKpi]:
        def pct(value: float) -> str:
            return f"{round(value * 100)}%"

        return [
            DraftMetricKpi(
                key="cycle_time",
                label="Avg cycle time",
                value=response.average_cycle_hours,
                formatted_value=f"{response.average_cycle_hours:.1f}h",
                target="< 48h",
                status="good" if response.average_cycle_hours <= 48 else "watch",
                detail="Average time from draft run start to latest completed lifecycle point.",
            ),
            DraftMetricKpi(
                key="first_review_approval",
                label="First-review approval",
                value=response.first_review_approval_rate,
                formatted_value=pct(response.first_review_approval_rate),
                target="> 75%",
                status="good" if response.first_review_approval_rate >= 0.75 else "watch",
                detail="Approved runs that were not returned for correction in the current window.",
            ),
            DraftMetricKpi(
                key="unsupported_claims",
                label="Unsupported claims",
                value=response.unsupported_claim_rate,
                formatted_value=pct(response.unsupported_claim_rate),
                target="< 2%",
                status="good" if response.unsupported_claim_rate <= 0.02 else "risk",
                detail="Draft runs with unsupported source or clause findings.",
            ),
            DraftMetricKpi(
                key="source_integrity",
                label="Source integrity",
                value=response.source_integrity_rate,
                formatted_value=pct(response.source_integrity_rate),
                target="> 95%",
                status="good" if response.source_integrity_rate >= 0.95 else "watch",
                detail="Runs with at least one source and source hashes on all source ledger entries.",
            ),
            DraftMetricKpi(
                key="artifact_compliance",
                label="Issued artifact compliance",
                value=response.issue_artifact_compliance_rate,
                formatted_value=pct(response.issue_artifact_compliance_rate),
                target="100%",
                status="good" if response.issue_artifact_compliance_rate >= 1 else "risk",
                detail="Issued letters with immutable DOCX and PDF artifact references.",
            ),
            DraftMetricKpi(
                key="overdue_reviews",
                label="Overdue reviews",
                value=float(response.overdue_review_count),
                formatted_value=str(response.overdue_review_count),
                target="0",
                status="good" if response.overdue_review_count == 0 else "risk",
                detail="Assigned review tasks past their due date.",
            ),
        ]

    async def _export_artifacts(
        self,
        letter_id: str,
        run: DraftRun,
        current_user: Any,
    ) -> Dict[str, Any]:
        letter = await self.letter_service.get_letter(letter_id)
        if not letter:
            raise HTTPException(status_code=404, detail="Letter not found")
        org_id = str(getattr(letter, "organization_id", "") or "")
        project_id = str(getattr(letter, "project_id", "") or "")
        if not org_id:
            raise HTTPException(status_code=400, detail="Letter organization is required for export")
        body = (run.draft_artifact.draft_letter if run.draft_artifact else "") or ""
        subject = str(getattr(letter, "subject", None) or run.inputs.get("subject") or "Letter Draft")
        safe_run = str(run.run_id).replace("/", "-")
        base_key = f"letter-drafts/{letter_id}/{safe_run}"
        file_service = FileObjectService(self.db)
        docx_result = await file_service.store_bytes(
            content=self._build_docx_bytes(subject, body),
            organization_id=org_id,
            project_id=project_id,
            original_filename=f"{safe_run}.docx",
            storage_key=f"{base_key}/{safe_run}.docx",
            current_user=current_user,
            document_type="letter_export",
            upload_id=run.run_id,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            document_id=letter_id,
        )
        pdf_result = await file_service.store_bytes(
            content=self._build_pdf_bytes(subject, body),
            organization_id=org_id,
            project_id=project_id,
            original_filename=f"{safe_run}.pdf",
            storage_key=f"{base_key}/{safe_run}.pdf",
            current_user=current_user,
            document_type="letter_export",
            upload_id=run.run_id,
            content_type="application/pdf",
            document_id=letter_id,
        )
        return {
            "docx_file_object_id": docx_result.get("file_object_id"),
            "pdf_file_object_id": pdf_result.get("file_object_id"),
            "docx_storage_key": docx_result.get("storage_key"),
            "pdf_storage_key": pdf_result.get("storage_key"),
        }

    @staticmethod
    def _build_docx_bytes(subject: str, body: str) -> bytes:
        try:
            from docx import Document  # type: ignore

            buffer = io.BytesIO()
            doc = Document()
            doc.add_heading(subject or "Letter Draft", level=1)
            for paragraph in (body or "").splitlines():
                doc.add_paragraph(paragraph)
            doc.save(buffer)
            return buffer.getvalue()
        except Exception:
            escaped_subject = (subject or "Letter Draft").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            escaped_body = (body or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            document_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body>"
                f"<w:p><w:r><w:t>{escaped_subject}</w:t></w:r></w:p>"
                f"<w:p><w:r><w:t>{escaped_body}</w:t></w:r></w:p>"
                "</w:body></w:document>"
            )
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(
                    "[Content_Types].xml",
                    (
                        '<?xml version="1.0" encoding="UTF-8"?>'
                        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                        '<Default Extension="xml" ContentType="application/xml"/>'
                        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                        "</Types>"
                    ),
                )
                archive.writestr(
                    "_rels/.rels",
                    (
                        '<?xml version="1.0" encoding="UTF-8"?>'
                        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                        "</Relationships>"
                    ),
                )
                archive.writestr("word/document.xml", document_xml)
            return buffer.getvalue()

    @staticmethod
    def _build_pdf_bytes(subject: str, body: str) -> bytes:
        try:
            from reportlab.lib.pagesizes import A4  # type: ignore
            from reportlab.pdfgen import canvas  # type: ignore

            buffer = io.BytesIO()
            pdf = canvas.Canvas(buffer, pagesize=A4)
            width, height = A4
            y = height - 72
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(72, y, (subject or "Letter Draft")[:100])
            y -= 28
            pdf.setFont("Helvetica", 10)
            for line in (body or "").splitlines():
                if y < 72:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 10)
                    y = height - 72
                pdf.drawString(72, y, line[:110])
                y -= 14
            pdf.save()
            return buffer.getvalue()
        except Exception:
            def escape_pdf(value: str) -> str:
                return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

            lines = [subject or "Letter Draft", *[line for line in (body or "").splitlines() if line.strip()]]
            text_ops = ["BT", "/F1 10 Tf", "72 760 Td"]
            for index, line in enumerate(lines[:45]):
                if index:
                    text_ops.append("0 -14 Td")
                text_ops.append(f"({escape_pdf(line[:100])}) Tj")
            text_ops.append("ET")
            stream = "\n".join(text_ops).encode("latin-1", errors="replace")
            objects = [
                b"<< /Type /Catalog /Pages 2 0 R >>",
                b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
                b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
                b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            ]
            buffer = io.BytesIO()
            buffer.write(b"%PDF-1.4\n")
            offsets = []
            for idx, obj in enumerate(objects, start=1):
                offsets.append(buffer.tell())
                buffer.write(f"{idx} 0 obj\n".encode("ascii"))
                buffer.write(obj)
                buffer.write(b"\nendobj\n")
            xref_at = buffer.tell()
            buffer.write(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
            buffer.write(b"0000000000 65535 f \n")
            for offset in offsets:
                buffer.write(f"{offset:010d} 00000 n \n".encode("ascii"))
            buffer.write(
                f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode(
                    "ascii"
                )
            )
            return buffer.getvalue()

    @staticmethod
    def _source_ledger_hash(sources: List[SourceEvidence]) -> str:
        payload = "|".join(sorted(source.source_hash or source.source_id for source in sources))
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    async def _emit_notification(
        self,
        event_type: NotificationType,
        letter_id: str,
        current_user: Any,
        *,
        include_users: Optional[List[str]],
        title: str,
        message: str,
    ) -> None:
        if not include_users:
            return
        try:
            letter = await self.letter_service.get_letter(letter_id)
            await NotificationService(self.db).emit(
                event_type,
                resource_id=str(letter_id),
                resource_type="letter",
                actor_id=self._user_id(current_user),
                include_users=include_users,
                priority=NotificationPriority.HIGH,
                severity=NotificationSeverity.INFO,
                data={
                    "title": title,
                    "message": message,
                    "organization_id": str(getattr(letter, "organization_id", "") or "") if letter else None,
                    "project_id": str(getattr(letter, "project_id", "") or "") if letter else None,
                },
            )
        except Exception:
            return

    def _guardrail_scan(
        self,
        request: DraftRunCreateRequest,
        sources: List[SourceEvidence],
    ) -> GuardrailReport:
        """Deterministic injection scan over user inputs and source texts.

        Critical findings block drafting and approval. Lower-severity findings
        remain visible to the three-stage human approval chain.
        """
        findings = []
        for origin in ("subject", "points", "context", "requirements", "background_facts"):
            findings.extend(self.guardrails.scan_input(getattr(request, origin, None), origin=origin))
        findings.extend(
            self.guardrails.scan_evidence([source.text or source.snippet for source in sources])
        )
        verdict = (
            "requires_human_review"
            if any(finding.severity == "critical" for finding in findings)
            else "pass"
        )
        return GuardrailReport(verdict=verdict, findings=findings)

    @staticmethod
    def _evidence_ledger_entries(
        run_id: str, sources: List[SourceEvidence]
    ) -> List[EvidenceLedgerEntry]:
        """Unified provenance records; labels match the prompt's [S#] tokens."""
        return [
            EvidenceLedgerEntry.from_source_evidence(
                source, run_id=run_id, citation_label=f"S{idx}"
            )
            for idx, source in enumerate(sources, start=1)
        ]

    async def _record_run_observability(self, run: DraftRun) -> None:
        """Count run outcomes and degraded/guardrail signals; never fails the run."""
        try:
            from ...services.observability import observability_registry

            await observability_registry.record_domain_event(
                resource_type="letter_drafting", event_type=f"run_{run.status}"
            )
            if any(str(w).startswith("draft_llm:") for w in run.warnings):
                await observability_registry.record_domain_event(
                    resource_type="letter_drafting", event_type="llm_fallback_draft"
                )
            if run.guardrail_report and run.guardrail_report.findings:
                await observability_registry.record_domain_event(
                    resource_type="letter_drafting", event_type="prompt_injection_suspected"
                )
        except Exception:
            # Observability must never break drafting; the run record itself
            # still carries the warnings and reports.
            pass

    async def _create_and_record(
        self,
        run: DraftRun,
        current_user: Any,
        *,
        context_pack: Optional[DraftContextPack] = None,
    ) -> DraftRun:
        await self._record_run_observability(run)
        stored = await self.repository.create(run)
        if stored.run_id != run.run_id:
            # A duplicate idempotent request won the creation race.  It owns
            # the snapshots and audit event; do not create duplicate evidence.
            return stored
        snapshot_fields = await self.repository.create_immutable_snapshots(stored)
        stored = await self.repository.update_fields(
            stored.letter_id,
            stored.run_id,
            snapshot_fields,
        ) or stored
        if context_pack:
            await self.repository.create_context_pack(context_pack)
        await self.repository.append_event(
            stored.letter_id,
            stored.run_id,
            "created",
            actor_user_id=self._user_id(current_user),
            status=stored.status,
            payload={"mode": stored.mode, "draft_type": stored.draft_type},
        )
        return stored

    async def _add_governance_comment_context(
        self,
        letter_id: str,
        context: DraftContextBundle,
        sources: List[SourceEvidence],
        *,
        current_run_id: str,
    ) -> List[SourceEvidence]:
        latest = await self.repository.latest(letter_id, None)
        if not latest or latest.run_id == current_run_id:
            return sources
        comments = await self.repository.list_comments(letter_id, latest.run_id)
        if not comments:
            return sources
        existing_lines = set(context.comments or [])
        enriched = list(sources)
        for index, comment in enumerate(comments[-10:], start=1):
            body = self._snippet(comment.body, width=800)
            if not body:
                continue
            if body not in existing_lines:
                context.comments.append(body)
                existing_lines.add(body)
            enriched.append(
                SourceEvidence(
                    source_id=f"governance-comment:{latest.run_id}:{comment.comment_id}",
                    source_type="comment",
                    allowed_use="comment",
                    organization_id=context.active_workspace.get("organization_id"),
                    project_id=context.active_workspace.get("project_id"),
                    label=f"Draft governance comment {index}",
                    text=body,
                    snippet=self._snippet(body, width=240),
                    letter_id=str(letter_id),
                    metadata={
                        "source_run_id": latest.run_id,
                        "comment_id": comment.comment_id,
                        "visibility": comment.visibility,
                        "created_by": comment.created_by,
                    },
                )
            )
        return enriched

    async def _run_cyclic_draft(
        self,
        *,
        letter: Letter,
        request: DraftRunCreateRequest,
        current_user: Any,
        role: str,
        recipient_focus: Optional[str],
        context: Any,
        sources: List[SourceEvidence],
        inputs: Dict[str, Any],
        plan: str,
        generator: DraftGenerator,
        initial_artifact: DraftArtifact,
    ) -> tuple[
        DraftArtifact,
        List[SourceEvidence],
        ValidationReport,
        List[CyclicIterationTrace],
        List[DraftAssertionSupport],
        DraftConfidenceScores,
        List[str],
    ]:
        artifact = initial_artifact
        working_sources = self._with_source_hashes(sources)
        trace: List[CyclicIterationTrace] = []
        warnings: List[str] = []
        validation = self._merge_reports(
            self.validator.critique(
                artifact,
                role,
                working_sources,
                finalized=request.finalized,
            ),
            self.validator.strategy_alignment(artifact, plan),
        )
        max_iterations = max(1, min(int(request.max_iterations or 3), 5))

        for iteration in range(1, max_iterations + 1):
            refinement_queries = self._refinement_queries(validation)
            retrieved: List[SourceEvidence] = []
            regenerated = False

            if validation.blocking and refinement_queries:
                retrieved = await self._retrieve_refinement_sources(
                    refinement_queries,
                    context,
                    current_user,
                )
                if retrieved:
                    known_ids = {source.source_id for source in working_sources}
                    additions = [source for source in retrieved if source.source_id not in known_ids]
                    if additions:
                        working_sources = self._with_source_hashes(working_sources + additions)
                        artifact, draft_warnings = await generator.generate(
                            letter,
                            role,
                            recipient_focus,
                            context,
                            working_sources,
                            inputs,
                            plan=plan,
                            finalized=request.finalized,
                        )
                        warnings.extend(draft_warnings)
                        regenerated = True
                        validation = self._merge_reports(
                            self.validator.critique(
                                artifact,
                                role,
                                working_sources,
                                finalized=request.finalized,
                            ),
                            self.validator.strategy_alignment(artifact, plan),
                        )

            trace.append(
                CyclicIterationTrace(
                    iteration=iteration,
                    critique_blocking=validation.blocking,
                    finding_codes=[finding.code for finding in validation.findings],
                    refinement_queries=refinement_queries,
                    retrieved_source_ids=[source.source_id for source in retrieved],
                    regenerated=regenerated,
                    notes=(
                        "Stopping conditions satisfied."
                        if not validation.blocking
                        else "Blocking findings remain."
                    ),
                )
            )

            if not validation.blocking:
                break
            if not retrieved or not regenerated:
                break

        assertion_support = self._assertion_support(artifact, working_sources, validation)
        confidence_scores = self._confidence_scores(validation, assertion_support, working_sources)
        return (
            artifact,
            working_sources,
            validation,
            trace,
            assertion_support,
            confidence_scores,
            warnings,
        )

    async def _retrieve_refinement_sources(
        self,
        refinement_queries: List[str],
        context: Any,
        current_user: Any,
    ) -> List[SourceEvidence]:
        org_id = context.active_workspace.get("organization_id")
        project_id = context.active_workspace.get("project_id")
        if not org_id or not project_id:
            return []
        retrieved: List[SourceEvidence] = []
        for query in refinement_queries:
            if not query:
                continue
            try:
                result = await self.exact_clause_search(
                    ExactClauseSearchRequest(
                        organization_id=str(org_id),
                        project_id=str(project_id),
                        clause_number=query,
                        limit=5,
                    ),
                    current_user,
                )
                retrieved.extend(result.sources)
            except Exception:
                continue
        return self._with_source_hashes(retrieved)

    @staticmethod
    def _refinement_queries(validation: ValidationReport) -> List[str]:
        queries: List[str] = []
        for finding in validation.findings:
            if finding.code != "unsupported_clause_citation" or not finding.evidence:
                continue
            for value in finding.evidence.split(","):
                cleaned = value.strip()
                if cleaned and cleaned not in queries:
                    queries.append(cleaned)
        return queries[:5]

    def _assertion_support(
        self,
        artifact: DraftArtifact,
        sources: List[SourceEvidence],
        validation: ValidationReport,
    ) -> List[DraftAssertionSupport]:
        sentences = [
            sentence.strip()
            for sentence in re.split(r"(?<=[.!?])\s+", artifact.draft_letter or "")
            if len(sentence.strip()) >= 30
        ]
        supports: List[DraftAssertionSupport] = []
        for sentence in sentences[:20]:
            source_ids = self._matching_source_ids(sentence, sources)
            if "[CONFIRM:" in sentence or "[TO BE INSERTED BY USER:" in sentence:
                status = "needs_confirmation"
                risk = "medium"
            elif source_ids:
                status = "supported"
                risk = "low"
            elif any(source.source_type == "current_input" for source in sources):
                status = "user_provided"
                risk = "medium"
            else:
                status = "unsupported"
                risk = "high"
            supports.append(
                DraftAssertionSupport(
                    assertion_id=str(uuid.uuid4()),
                    text=sentence,
                    support_status=status,
                    source_ids=source_ids,
                    risk_level=risk,
                )
            )
        if validation.blocking and not supports:
            supports.append(
                DraftAssertionSupport(
                    assertion_id=str(uuid.uuid4()),
                    text="Draft validation contains blocking findings.",
                    support_status="unsupported",
                    source_ids=[],
                    risk_level="high",
                )
            )
        return supports

    @staticmethod
    def _matching_source_ids(sentence: str, sources: List[SourceEvidence]) -> List[str]:
        terms = {
            term.lower()
            for term in re.findall(r"[A-Za-z0-9][A-Za-z0-9._-]{3,}", sentence)
            if term.lower() not in {"contractor", "engineer", "employer", "shall", "this", "that", "with"}
        }
        if not terms:
            return []
        matches: List[str] = []
        for source in sources:
            haystack = " ".join(
                str(value or "")
                for value in [
                    source.label,
                    source.text,
                    source.snippet,
                    source.clause_number,
                    source.clause_title,
                ]
            ).lower()
            hit_count = sum(1 for term in terms if term in haystack)
            if hit_count >= max(2, min(4, len(terms) // 4)):
                matches.append(source.source_id)
        return matches[:5]

    @staticmethod
    def _confidence_scores(
        validation: ValidationReport,
        assertions: List[DraftAssertionSupport],
        sources: List[SourceEvidence],
    ) -> DraftConfidenceScores:
        error_count = sum(1 for finding in validation.findings if finding.level == "error")
        warning_count = sum(1 for finding in validation.findings if finding.level == "warning")
        clause_sources = [source for source in sources if source.source_type == "contract_clause"]
        clause_confidence = 1.0 if clause_sources else 0.55
        if any(finding.code == "unsupported_clause_citation" for finding in validation.findings):
            clause_confidence = 0.25

        if assertions:
            supported = sum(
                1
                for assertion in assertions
                if assertion.support_status in {"supported", "user_provided"}
            )
            factual_support = supported / len(assertions)
        else:
            factual_support = 0.5

        tone_suitability = max(0.0, 1.0 - (warning_count * 0.08) - (error_count * 0.2))
        overall = max(
            0.0,
            min(
                1.0,
                (clause_confidence * 0.35)
                + (factual_support * 0.4)
                + (tone_suitability * 0.25)
                - (error_count * 0.12),
            ),
        )
        risk_level = "high" if error_count else ("medium" if warning_count or overall < 0.85 else "low")
        return DraftConfidenceScores(
            clause_confidence=round(clause_confidence, 3),
            factual_support=round(factual_support, 3),
            tone_suitability=round(tone_suitability, 3),
            overall=round(overall, 3),
            risk_level=risk_level,
        )

    def _build_context_pack(
        self,
        *,
        letter: Letter,
        request: DraftRunCreateRequest,
        run_id: str,
        context: Any,
        sources: List[SourceEvidence],
        source_warnings: List[str],
    ) -> DraftContextPack:
        inputs = self._inputs_payload(letter, request)
        contractual_basis = [
            {
                "source_id": source.source_id,
                "clause_number": source.clause_number,
                "clause_title": source.clause_title,
                "label": source.label,
                "snippet": source.snippet,
                "document_id": source.document_id,
                "page_numbers": source.page_numbers,
            }
            for source in sources
            if source.source_type == "contract_clause"
        ]
        prior_correspondence = [
            {
                "source_id": source.source_id,
                "letter_id": source.letter_id,
                "label": source.label,
                "snippet": source.snippet,
            }
            for source in sources
            if source.source_type in {"prior_correspondence", "graph_thread"}
        ]
        missing_confirmations = [
            key.replace("_", " ")
            for key, present in (context.threshold_inputs or {}).items()
            if not present
        ]
        required_actions = [
            value
            for value in [
                request.required_action,
                request.desired_position,
                request.purpose,
            ]
            if value
        ]
        facts = list(context.current_materials or [])[:20]
        return DraftContextPack(
            context_pack_id=str(uuid.uuid4()),
            letter_id=str(getattr(letter, "id", "") or ""),
            run_id=run_id,
            project={
                "organization_id": context.active_workspace.get("organization_id"),
                "project_id": context.active_workspace.get("project_id"),
                "letter_id": context.active_workspace.get("letter_id"),
                "letter_no": context.active_workspace.get("letter_no"),
            },
            draft_request={
                "draft_type": request.draft_type,
                "letter_category": request.letter_category,
                "contract_package": request.contract_package,
                "role": request.role or getattr(letter, "strategy_role", None),
                "subject": inputs.get("subject"),
                "recipient": inputs.get("recipient"),
                "tone": request.tone,
                "trigger_event": request.trigger_event,
            },
            facts=facts,
            contractual_basis=contractual_basis,
            prior_correspondence=prior_correspondence,
            required_actions=required_actions,
            risk_flags=[],
            missing_confirmations=missing_confirmations,
            source_notes=list(source_warnings or []),
            source_ids=[source.source_id for source in sources],
        )

    @staticmethod
    def _assert_scope_allowed(current_user: Any, organization_id: str, project_id: str) -> None:
        roles = {str(role).lower() for role in (getattr(current_user, "roles", []) or [])}
        if "superadmin" in roles:
            return
        user_org = str(getattr(current_user, "organization_id", "") or "")
        if user_org and str(organization_id) != user_org:
            raise HTTPException(status_code=403, detail="Not authorized for this organization")
        allowed_projects = {
            str(project)
            for project in (getattr(current_user, "projects", []) or [])
            if project
        }
        if allowed_projects and str(project_id) not in allowed_projects:
            raise HTTPException(status_code=403, detail="Not authorized for this project")

    @staticmethod
    def _snippet(value: Any, width: int = 500) -> Optional[str]:
        if not value:
            return None
        text = " ".join(str(value).split())
        if not text:
            return None
        if len(text) <= width:
            return text
        return f"{text[: max(0, width - 3)]}..."

    @staticmethod
    def _with_source_hashes(sources: List[SourceEvidence]) -> List[SourceEvidence]:
        normalized: List[SourceEvidence] = []
        for source in sources:
            if source.source_hash:
                normalized.append(source)
                continue
            raw = "|".join(
                [
                    source.source_id,
                    source.source_type,
                    source.allowed_use,
                    source.document_id or "",
                    source.letter_id or "",
                    source.clause_number or "",
                    source.text or source.snippet or "",
                ]
            )
            normalized.append(
                source.model_copy(
                    update={
                        "source_hash": hashlib.sha256(raw.encode("utf-8")).hexdigest()
                    }
                )
            )
        return normalized

    async def _load_and_authorize(
        self,
        letter_id: str,
        current_user: Any,
        action: str,
        *,
        drafting_permission: Optional[str] = None,
    ) -> Letter:
        letter = await self.letter_service.get_letter(letter_id)
        if not letter:
            raise HTTPException(status_code=404, detail="Letter not found")
        if drafting_permission:
            await self.policy_service.authorize(
                current_user,
                drafting_permission,
                resource_type="letter",
                resource_id=str(letter_id),
                organization_id=str(getattr(letter, "organization_id", "") or ""),
                project_id=str(getattr(letter, "project_id", "") or ""),
                letter_id=str(letter_id),
            )
            return letter
        await self.auth_service.check_letter_access(current_user, letter, action)
        return letter

    @staticmethod
    def _resolve_role(letter: Letter, request: DraftRunCreateRequest) -> DraftRole:
        raw = (request.role or getattr(letter, "strategy_role", None) or "contractor").lower()
        if raw.startswith("engineer"):
            return "engineer"
        if raw.startswith("employer"):
            return "employer"
        return "contractor"

    @staticmethod
    def _inputs_payload(letter: Letter, request: DraftRunCreateRequest) -> Dict[str, Any]:
        return {
            "draft_type": request.draft_type,
            "letter_category": request.letter_category,
            "contract_package": request.contract_package,
            "role": request.role or getattr(letter, "strategy_role", None),
            "subject": request.subject or letter.subject,
            "recipient": request.recipient or letter.recipient,
            "requirements": request.requirements,
            "points": request.points,
            "purpose": request.purpose,
            "desired_position": request.desired_position,
            "required_action": request.required_action,
            "background_facts": request.background_facts,
            "trigger_event": request.trigger_event,
            "tone": request.tone,
            "timeline_days": request.timeline_days,
            "incoming_document_id": request.incoming_document_id,
            "incoming_letter_id": request.incoming_letter_id,
            "clauses_to_consider": list(request.clauses_to_consider),
            "attachments": list(request.attachments),
            "document_ids": list(request.document_ids),
            "include_letter_codes": list(request.include_letter_codes),
            "exclude_letter_codes": list(request.exclude_letter_codes),
            "plan_override": request.plan_override,
            "max_iterations": request.max_iterations,
            "finalized": request.finalized,
            "user_direction_answers": [
                answer.model_dump(mode="json") for answer in request.user_direction_answers
            ],
            "user_direction": request.user_direction,
        }

    @staticmethod
    def _blocked_artifact(report: ValidationReport) -> DraftArtifact:
        required = ", ".join(finding.message for finding in report.findings)
        return DraftArtifact(
            draft_letter="",
            source_integrity_notes=(
                "DRAFT BLOCKED: Insufficient source material. Required: "
                f"{required}"
            ),
            raw_model_output="",
        )

    @staticmethod
    def _merge_reports(*reports: ValidationReport) -> ValidationReport:
        findings = []
        for report in reports:
            findings.extend(report.findings)
        return ValidationReport(
            blocking=any(finding.level == "error" for finding in findings),
            findings=findings,
        )

    async def _resolve_strategy_plan(
        self,
        letter_id: str,
        letter: Letter,
        request: DraftRunCreateRequest,
    ) -> str:
        if request.mode == "strategy":
            return ""
        if request.mode in _AI_DRAFT_MODES and not self._strategy_is_approved(letter):
            return ""
        if request.plan_override:
            return request.plan_override
        saved_plan = getattr(letter, "strategy_plan", None) or getattr(letter, "draft_plan", None)
        if saved_plan:
            return saved_plan
        latest_strategy = await self.repository.latest(letter_id, "strategy")
        if latest_strategy and latest_strategy.plan and latest_strategy.status != "failed":
            return latest_strategy.plan
        return ""

    @staticmethod
    def _strategy_is_approved(letter: Letter) -> bool:
        if getattr(letter, "strategy_plan_approved_at", None):
            return True
        if getattr(letter, "accepted_strategy_version", None):
            return True
        # Backward compatibility for letters already advanced before the
        # approval audit fields existed.
        return str(getattr(letter, "status", "") or "").lower() in {
            "draft",
            "review",
            "approval",
            "completed",
        }

    @staticmethod
    def _revision_instruction(request: ReviseDraftRequest) -> str:
        labels = {
            "make_firmer": "Revise the draft to be firmer while staying professional and source-faithful.",
            "make_more_polite": "Revise the draft to be more polite and conciliatory without conceding unsupported points.",
            "add_contractual_reasoning": "Add clearer contractual reasoning using only available clause evidence.",
            "add_clause_reference": "Add clause references only where the retrieved source ledger supports them.",
            "make_short": "Shorten the draft while preserving the required contractual position.",
            "make_detailed": "Expand the draft with more structured facts, reasoning, and source integrity notes.",
            "convert_to_employer_submission": "Revise from an Employer profile if supported by the request context.",
            "convert_to_contractor_letter": "Revise from a Contractor profile if supported by the request context.",
            "regenerate": "Regenerate the draft from the confirmed plan and source ledger.",
            "custom_instruction": request.custom_instruction or "Apply the user's custom revision instruction.",
        }
        instruction = labels.get(request.revision_action, request.revision_action)
        if request.additional_requirements:
            instruction = f"{instruction}\nAdditional requirements: {request.additional_requirements}"
        return instruction

    @staticmethod
    def _section_integrity_violations(run: DraftRun) -> List[int]:
        if not run.draft_artifact:
            return []
        checks: List[FrozenDraftSection] = list(run.frozen_sections or [])
        if run.section_revision:
            checks.extend(run.section_revision.preserved_sections)
        return sorted(set(verify_frozen_sections(run.draft_artifact.draft_letter, checks)))

    @staticmethod
    def _user_id(current_user: Any) -> Optional[str]:
        return (
            getattr(current_user, "id", None)
            or getattr(current_user, "email", None)
            or getattr(current_user, "username", None)
        )
